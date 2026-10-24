import os
import datetime
import json
import io
import re
from flask import request, redirect, url_for, flash, jsonify, abort, send_file, current_app
from flask_login import login_user, login_required, logout_user, current_user
from markupsafe import escape
from app.utils.notifications import notify_student_and_parent, notify_users
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from app.extensions import db
from app.models import User, Attendance, PendingCard, FeeTransaction, Book, DisciplineCase, SchoolMessage, Grade, TimetableEntry, DayScheduleEntry, Setting, AcademicSubject, SyllabusCoverage, LessonPlan, ReportCardApproval
from app.utils.constants import *
from app.utils.decorators import role_required
from app.utils.helpers import get_setting, set_setting, attendance_for_student, linked_children, teacher_assignments, teacher_assigned_classes, get_school_logo_static_url
from app.utils.rendering import page
from app.utils.decorators import permission_for

def timetable():
    if request.method == "POST" and current_user.role not in {"master_admin", "secretary", "patron", "matron"}:
        abort(403)
    if request.method == "POST":
        class_name = request.form.get("class_name", "").strip()
        day_of_week = request.form.get("day_of_week", "").strip()
        time_slot = request.form.get("time_slot", "").strip()
        subject = request.form.get("subject", "").strip()
        teacher_name = request.form.get("teacher_name", "").strip()
        assigned_teacher = User.query.filter_by(role="teacher", full_name=teacher_name, class_name=class_name).first() if teacher_name else None
        if not class_name or not day_of_week or not subject:
            flash("Class name, day, and subject are required.", "danger")
        elif teacher_name and not assigned_teacher:
            flash("Choose a teacher assigned to this class.", "danger")
        elif current_user.role in {"patron", "matron"} and class_name != current_user.class_name:
            abort(403)
        else:
            t = TimetableEntry(class_name=class_name, day_of_week=day_of_week, time_slot=time_slot, subject=subject, teacher_name=teacher_name)
            db.session.add(t)
            db.session.commit()
            flash("Timetable entry added.", "success")
            return redirect(url_for("timetable"))
    entry_query = TimetableEntry.query
    if current_user.role == "teacher":
        entry_query = entry_query.filter(TimetableEntry.class_name.in_(teacher_assigned_classes(current_user)))
    elif current_user.role in {"discipline", "patron", "matron"}:
        entry_query = entry_query.filter_by(class_name=current_user.class_name)
    elif current_user.role == "student":
        entry_query = entry_query.filter_by(class_name=current_user.class_name)
    elif current_user.role == "parent":
        child_classes = {child.class_name for child in linked_children(current_user.id) if child.class_name}
        entry_query = entry_query.filter(TimetableEntry.class_name.in_(child_classes))
    entries = entry_query.order_by(TimetableEntry.class_name, TimetableEntry.day_of_week, TimetableEntry.time_slot).all()
    rows = "".join(f"<tr><td>{e.class_name}</td><td>{e.day_of_week}</td><td>{e.time_slot}</td><td>{e.subject}</td><td>{e.teacher_name or '-'}</td></tr>" for e in entries)
    can_manage = current_user.role in {"master_admin", "secretary", "patron", "matron"}
    teacher_query = User.query.filter_by(role="teacher")
    if current_user.role in {"patron", "matron"}:
        teacher_query = teacher_query.filter_by(class_name=current_user.class_name)
    teacher_options = "".join(f"<option value='{escape(t.full_name)}'>{escape(t.full_name)}</option>" for t in teacher_query.order_by(User.full_name).all())
    if current_user.role == "parent":
        intro = "Class schedules for your linked children."
    elif current_user.role == "student":
        intro = "Your academic timetable."
    else:
        intro = "Class schedules and lesson times."
    timetable_form = f"""<form method="post" class="row g-3"><div class="col-md-4"><label>Class Name</label><input class="form-control" name="class_name" required></div><div class="col-md-4"><label>Day of Week</label><select class="form-select" name="day_of_week"><option>Monday</option><option>Tuesday</option><option>Wednesday</option><option>Thursday</option><option>Friday</option></select></div><div class="col-md-4"><label>Time Slot</label><input class="form-control" name="time_slot" placeholder="e.g. 08:00 - 09:20"></div><div class="col-md-6"><label>Subject / Lesson</label><input class="form-control" name="subject" required></div><div class="col-md-6"><label>Assigned Teacher</label><select class="form-select" name="teacher_name"><option value="">Choose teacher...</option>{teacher_options}</select></div><div class="col-12"><button class="btn btn-accent">Add Timetable Entry</button></div></form>""" if can_manage else ""
    body = f"""<section class="panel"><h1>Class <span class="accent">Timetable</span></h1><p class="muted">{intro}</p>{timetable_form}</section>
<section class="panel"><h2>{'Existing Timetable' if can_manage else 'Academic Timetable'}</h2><div class="table-responsive"><table class="table"><tr><th>Class</th><th>Day</th><th>Time Slot</th><th>Subject</th><th>Teacher</th></tr>{rows or '<tr><td colspan="5" class="muted">No timetable has been published for the relevant class yet.</td></tr>'}</table></div></section>"""
    return page("Timetable", body, "teacher")

def day_schedule():
    managers = {"master_admin", "secretary", "patron", "matron"}
    can_manage = current_user.role in managers
    if request.method == "POST":
        if not can_manage:
            abort(403)
        action = request.form.get("action", "create")
        if action == "set_publication":
            entry = db.session.get(DayScheduleEntry, request.form.get("entry_id", type=int)) or abort(404)
            if current_user.role in {"patron", "matron"} and (not current_user.class_name or entry.class_name != current_user.class_name):
                abort(403)
            entry.is_published = request.form.get("is_published") == "yes"
            db.session.commit()
            flash("Day schedule publication updated.", "success")
            return redirect(url_for("day_schedule"))

        raw_date = request.form.get("schedule_date", "").strip()
        class_name = request.form.get("class_name", "").strip() or None
        activity = request.form.get("activity", "").strip()
        start_time = request.form.get("start_time", "").strip()
        end_time = request.form.get("end_time", "").strip()
        details = request.form.get("details", "").strip()
        try:
            schedule_date = datetime.date.fromisoformat(raw_date)
        except ValueError:
            schedule_date = None
        if not schedule_date or not activity:
            flash("A valid date and activity are required.", "danger")
        elif current_user.role in {"patron", "matron"} and (not current_user.class_name or class_name != current_user.class_name):
            abort(403)
        else:
            entry = DayScheduleEntry(
                schedule_date=schedule_date, class_name=class_name,
                start_time=start_time or None, end_time=end_time or None,
                activity=activity, details=details or None,
                is_published=request.form.get("is_published") == "yes",
                created_by=current_user.id,
            )
            db.session.add(entry)
            db.session.commit()
            flash("Day schedule activity saved.", "success")
            return redirect(url_for("day_schedule"))

    entry_query = DayScheduleEntry.query
    if not can_manage:
        entry_query = entry_query.filter(DayScheduleEntry.is_published.is_(True), DayScheduleEntry.schedule_date >= datetime.date.today())
    elif current_user.role in {"patron", "matron"}:
        entry_query = entry_query.filter(db.or_(DayScheduleEntry.is_published.is_(True), DayScheduleEntry.class_name == current_user.class_name))
    if current_user.role in {"student", "teacher", "discipline", "patron", "matron"}:
        class_names = teacher_assigned_classes(current_user) if current_user.role == "teacher" else ([current_user.class_name] if current_user.class_name else [])
        entry_query = entry_query.filter(db.or_(DayScheduleEntry.class_name.is_(None), DayScheduleEntry.class_name.in_(class_names)))
    elif current_user.role == "parent":
        class_names = {child.class_name for child in linked_children(current_user.id) if child.class_name}
        entry_query = entry_query.filter(db.or_(DayScheduleEntry.class_name.is_(None), DayScheduleEntry.class_name.in_(class_names)))
    entries = entry_query.order_by(DayScheduleEntry.schedule_date, DayScheduleEntry.start_time).all()
    rows = []
    for entry in entries:
        time_label = " - ".join(value for value in (entry.start_time, entry.end_time) if value) or "-"
        scope = entry.class_name or "All school"
        details = f"<br><span class='muted'>{escape(entry.details)}</span>" if entry.details else ""
        publication = "Published" if entry.is_published else "Draft"
        manage_cell = ""
        if can_manage:
            checked = "checked" if entry.is_published else ""
            manage_cell = f"<td><form method='post' class='d-flex align-items-center gap-2'><input type='hidden' name='action' value='set_publication'><input type='hidden' name='entry_id' value='{entry.id}'><label class='small'><input type='checkbox' name='is_published' value='yes' {checked}> Published</label><button class='btn btn-sm btn-outline-success'>Save</button></form></td>"
        else:
            manage_cell = f"<td>{publication}</td>"
        rows.append(f"<tr><td>{entry.schedule_date.isoformat()}</td><td>{escape(scope)}</td><td>{escape(time_label)}</td><td><strong>{escape(entry.activity)}</strong>{details}</td>{manage_cell}</tr>")

    if current_user.role == "parent":
        intro = "Published school activities for your linked childrenâ€™s classes, plus school-wide events."
    elif current_user.role == "student":
        intro = "Published school activities for your class, plus school-wide events."
    else:
        intro = "Plan and publish date-specific activities for a class or the whole school."
    class_field = "<div class='col-md-6'><label>Class (leave blank for all school)</label><input class='form-control' name='class_name' placeholder='e.g. L3 SWD'></div>"
    if current_user.role in {"patron", "matron"}:
        class_field = f"<input type='hidden' name='class_name' value='{escape(current_user.class_name or '')}'><div class='col-md-6'><label>Assigned class</label><input class='form-control' value='{escape(current_user.class_name or 'No class assigned')}' disabled></div>"
    form = f"""<section class='panel'><h1>Day <span class='accent'>Schedule</span></h1><p class='muted'>{intro}</p>{f'''<form method="post" class="row g-3"><div class="col-md-6"><label>Date</label><input class="form-control" type="date" name="schedule_date" value="{datetime.date.today().isoformat()}" required></div>{class_field}<div class="col-md-6"><label>Start Time</label><input class="form-control" type="time" name="start_time"></div><div class="col-md-6"><label>End Time</label><input class="form-control" type="time" name="end_time"></div><div class="col-12"><label>Activity</label><input class="form-control" name="activity" required></div><div class="col-12"><label>Details</label><textarea class="form-control" name="details" rows="2"></textarea></div><div class="col-12"><label><input type="checkbox" name="is_published" value="yes"> Publish now</label></div><div class="col-12"><button class="btn btn-accent">Save Activity</button></div></form>''' if can_manage else ''}</section>"""
    rows_html = "".join(rows) or f"<tr><td colspan='5' class='muted'>No published day schedule activities are available{ ' for your linked children' if current_user.role == 'parent' else ''}.</td></tr>"
    body = form + f"<section class='panel'><h2>{'Planned Activities' if can_manage else 'Published Activities'}</h2><div class='table-responsive'><table class='table day-schedule-table'><tr><th>Date</th><th>Class</th><th>Time</th><th>Activity</th><th>{'Manage' if can_manage else 'Status'}</th></tr>{rows_html}</table></div></section>"
    return page("Day Schedule", body, "teacher")

def assigned_subjects_for_teacher(teacher):
    return teacher_assignments(teacher)


def academic_management():
    managers = current_user.role in {"master_admin", "head_teacher"}
    if request.method == "POST":
        action = request.form.get("action", "")
        if action == "add_subject":
            if not managers:
                abort(403)
            trade = request.form.get("trade", "").strip()
            level = request.form.get("level", "").strip()
            class_name = request.form.get("class_name", "").strip()
            subject_name = request.form.get("subject", "").strip()
            teacher_id = request.form.get("teacher_id", type=int)
            year = request.form.get("academic_year", "").strip() or None
            if CLASS_TRADE_LEVEL.get(class_name) != (trade, level) or not subject_name:
                flash("Choose a valid trade, level, class, and subject.", "danger")
            elif teacher_id and not User.query.filter_by(id=teacher_id, role="teacher").first():
                flash("Choose a valid teacher account.", "danger")
            elif AcademicSubject.query.filter_by(class_name=class_name, subject=subject_name, academic_year=year).first():
                flash("That subject is already listed for this class and academic year.", "warning")
            else:
                db.session.add(AcademicSubject(trade=trade, level=level, class_name=class_name,
                    subject=subject_name, teacher_id=teacher_id, academic_year=year))
                db.session.commit()
                flash("Program subject added.", "success")
                return redirect(url_for("academic_management"))
        elif action == "submit_plan":
            subject = db.session.get(AcademicSubject, request.form.get("subject_id", type=int)) or abort(404)
            if subject.teacher_id != current_user.id:
                abort(403)
            try:
                lesson_date = datetime.date.fromisoformat(request.form.get("lesson_date", ""))
            except ValueError:
                lesson_date = None
            topic = request.form.get("topic", "").strip()
            objectives = request.form.get("objectives", "").strip()
            if not lesson_date or not topic or not objectives:
                flash("Lesson date, topic, and learning objectives are required.", "danger")
            else:
                db.session.add(LessonPlan(subject_id=subject.id, teacher_id=current_user.id,
                    lesson_date=lesson_date, topic=topic, objectives=objectives,
                    activities=request.form.get("activities", "").strip(),
                    resources=request.form.get("resources", "").strip()))
                db.session.commit()
                flash("Lesson plan submitted for review.", "success")
                return redirect(url_for("academic_management"))
        elif action == "update_coverage":
            subject = db.session.get(AcademicSubject, request.form.get("subject_id", type=int)) or abort(404)
            if subject.teacher_id != current_user.id:
                abort(403)
            unit = request.form.get("unit_title", "").strip()
            planned = request.form.get("planned_lessons", type=int)
            completed = request.form.get("completed_lessons", type=int)
            if not unit or planned is None or completed is None or planned < 1 or completed < 0 or completed > planned:
                flash("Enter a unit and valid lesson counts (completed cannot exceed planned).", "danger")
            else:
                coverage = SyllabusCoverage.query.filter_by(subject_id=subject.id, unit_title=unit).first()
                if coverage is None:
                    coverage = SyllabusCoverage(subject_id=subject.id, unit_title=unit, updated_by=current_user.id)
                    db.session.add(coverage)
                coverage.planned_lessons = planned
                coverage.completed_lessons = completed
                coverage.notes = request.form.get("notes", "").strip()
                coverage.updated_by = current_user.id
                db.session.commit()
                flash("Syllabus coverage saved.", "success")
                return redirect(url_for("academic_management"))
        elif action == "review_plan":
            if not managers:
                abort(403)
            plan = db.session.get(LessonPlan, request.form.get("plan_id", type=int)) or abort(404)
            decision = request.form.get("decision", "")
            if decision not in {"Reviewed", "Revision requested"}:
                abort(400)
            plan.status = decision
            plan.review_comment = request.form.get("review_comment", "").strip()
            plan.reviewed_by = current_user.id
            plan.reviewed_at = datetime.datetime.utcnow()
            db.session.commit()
            flash("Lesson plan review saved.", "success")
            return redirect(url_for("academic_management"))
        else:
            abort(400)

    subject_query = AcademicSubject.query.filter_by(is_active=True)
    if not managers:
        subject_query = subject_query.filter_by(teacher_id=current_user.id)
    subjects = subject_query.order_by(AcademicSubject.trade, AcademicSubject.level, AcademicSubject.class_name, AcademicSubject.subject).all()
    subject_rows = []
    for subject in subjects:
        coverage = subject.coverage
        planned = sum(item.planned_lessons for item in coverage)
        completed = sum(item.completed_lessons for item in coverage)
        percent = round(completed * 100 / planned) if planned else 0
        plans_pending = sum(plan.status == "Pending" for plan in subject.lesson_plans)
        subject_rows.append(f"<tr><td>{escape(subject.trade)}</td><td>{escape(subject.level)}</td><td>{escape(subject.class_name)}</td><td>{escape(subject.subject)}</td><td>{escape(subject.teacher.full_name if subject.teacher else 'Unassigned')}</td><td>{completed}/{planned} ({percent}%)</td><td>{plans_pending} pending</td></tr>")
    subject_select = "".join(f"<option value='{s.id}'>{escape(s.class_name)} Â· {escape(s.subject)}</option>" for s in subjects if s.teacher_id == current_user.id)
    catalog_form = ""
    if managers:
        class_options = "".join(f"<option value='{escape(name)}'>{escape(name)}</option>" for name in CLASS_NAMES)
        teacher_options = "".join(f"<option value='{t.id}'>{escape(t.full_name)} Â· {escape(t.class_name or 'No class')}</option>" for t in User.query.filter_by(role="teacher").order_by(User.full_name).all())
        catalog_form = f"""<section class='panel'><h2>Programs and Subjects</h2><p class='muted'>Add a subject to a program class and assign any teacher, including one teaching across departments.</p><form method='post' class='row g-3'><input type='hidden' name='action' value='add_subject'><div class='col-md-3'><label>Trade / Department</label><select class='form-select' name='trade' id='academic-trade' required><option value=''>Choose trade</option>{''.join(f'<option>{escape(t)}</option>' for t in TRADES)}</select></div><div class='col-md-2'><label>Level</label><select class='form-select' name='level' id='academic-level' required><option value=''>Choose level</option></select></div><div class='col-md-4'><label>Class</label><select class='form-select' name='class_name' id='academic-class' required><option value=''>Choose class</option></select></div><div class='col-md-3'><label>Subject</label><input class='form-control' name='subject' required maxlength='160'></div><div class='col-md-4'><label>Assigned Teacher</label><select class='form-select' name='teacher_id'><option value=''>Unassigned</option>{teacher_options}</select></div><div class='col-md-4'><label>Academic Year (optional)</label><input class='form-control' name='academic_year' placeholder='2026-2027'></div><div class='col-12'><button class='btn btn-accent'>Add Program Subject</button></div></form></section>
<script>const academicMap={json.dumps({trade: levels for trade, levels in PROGRAM_LEVELS.items()})};const classMap={json.dumps(CLASS_GROUPS)};const t=document.getElementById('academic-trade'),l=document.getElementById('academic-level'),c=document.getElementById('academic-class');function refreshAcademicOptions(){{l.innerHTML='<option value="">Choose level</option>';c.innerHTML='<option value="">Choose class</option>';for(const level of (academicMap[t.value]||[]))l.add(new Option(level,level));}}t.addEventListener('change',refreshAcademicOptions);l.addEventListener('change',()=>{{c.innerHTML='<option value="">Choose class</option>';for(const name of (((classMap[t.value]||{{}})[l.value])||[]))c.add(new Option(name,name));}});</script>"""
    teacher_forms = ""
    if not managers:
        teacher_forms = f"""<section class='panel'><h2>Submit a Lesson Plan</h2><form method='post' class='row g-3'><input type='hidden' name='action' value='submit_plan'><div class='col-md-6'><label>Assigned Class and Subject</label><select class='form-select' name='subject_id' required>{subject_select}</select></div><div class='col-md-6'><label>Lesson Date</label><input class='form-control' type='date' name='lesson_date' required></div><div class='col-md-6'><label>Topic</label><input class='form-control' name='topic' required></div><div class='col-md-6'><label>Learning Objectives</label><textarea class='form-control' name='objectives' required rows='2'></textarea></div><div class='col-md-6'><label>Teaching Activities</label><textarea class='form-control' name='activities' rows='2'></textarea></div><div class='col-md-6'><label>Resources</label><textarea class='form-control' name='resources' rows='2'></textarea></div><div class='col-12'><button class='btn btn-accent'>Submit for Review</button></div></form></section>
<section class='panel'><h2>Update Syllabus Coverage</h2><form method='post' class='row g-3'><input type='hidden' name='action' value='update_coverage'><div class='col-md-5'><label>Assigned Class and Subject</label><select class='form-select' name='subject_id' required>{subject_select}</select></div><div class='col-md-3'><label>Unit / Topic</label><input class='form-control' name='unit_title' required></div><div class='col-md-2'><label>Planned Lessons</label><input class='form-control' type='number' name='planned_lessons' min='1' required></div><div class='col-md-2'><label>Completed Lessons</label><input class='form-control' type='number' name='completed_lessons' min='0' required></div><div class='col-12'><label>Notes</label><input class='form-control' name='notes'></div><div class='col-12'><button class='btn btn-accent'>Save Coverage</button></div></form></section>"""
    plans_query = LessonPlan.query.join(AcademicSubject)
    if not managers:
        plans_query = plans_query.filter(LessonPlan.teacher_id == current_user.id)
    plans = plans_query.order_by(LessonPlan.lesson_date.desc(), LessonPlan.submitted_at.desc()).limit(100).all()
    plan_rows = []
    for plan in plans:
        review = ""
        if managers:
            review = f"<td><form method='post' class='d-flex flex-column gap-2'><input type='hidden' name='action' value='review_plan'><input type='hidden' name='plan_id' value='{plan.id}'><select class='form-select form-select-sm' name='decision'><option>Reviewed</option><option>Revision requested</option></select><input class='form-control form-control-sm' name='review_comment' placeholder='Review comments'><button class='btn btn-sm btn-outline-success'>Save review</button></form></td>"
        else:
            review = f"<td>{escape(plan.status)}{f'<br><small>{escape(plan.review_comment)}</small>' if plan.review_comment else ''}</td>"
        plan_rows.append(f"<tr><td>{escape(plan.subject.class_name)}</td><td>{escape(plan.subject.subject)}</td><td>{escape(plan.teacher.full_name)}</td><td>{plan.lesson_date.isoformat()}</td><td>{escape(plan.topic)}<br><small>{escape(plan.objectives)}</small></td><td>{escape(plan.status)}</td>{review}</tr>")
    body = f"""<section class='panel'><h1>Academic <span class='accent'>Management</span></h1><p class='muted'>Manage program subjects, track syllabus coverage, review lesson plans, and monitor teaching and learning progress.</p></section>{catalog_form}{teacher_forms}
<section class='panel'><h2>Teaching and Learning Progress</h2><div class='table-responsive'><table class='table align-middle'><thead><tr><th>Trade</th><th>Level</th><th>Class</th><th>Subject</th><th>Teacher</th><th>Syllabus Coverage</th><th>Lesson Plans</th></tr></thead><tbody>{''.join(subject_rows) or '<tr><td colspan="7" class="muted">No program subjects have been assigned yet.</td></tr>'}</tbody></table></div></section>
<section class='panel'><h2>{'Teachersâ€™ Lesson Plans' if managers else 'My Lesson Plans'}</h2><div class='table-responsive'><table class='table align-middle'><thead><tr><th>Class</th><th>Subject</th><th>Teacher</th><th>Date</th><th>Topic and Objectives</th><th>Status</th><th>{'Review' if managers else 'Feedback'}</th></tr></thead><tbody>{''.join(plan_rows) or '<tr><td colspan="7" class="muted">No lesson plans have been submitted.</td></tr>'}</tbody></table></div></section>"""
    return page("Academic Management", body, "teacher")


def grades():
    staff_roles = {"master_admin", "head_teacher", "teacher"}
    assigned_subjects = assigned_subjects_for_teacher(current_user) if current_user.role == "teacher" else set()
    assigned_subject_keys = {(class_name.casefold(), subject.casefold()) for class_name, subject in assigned_subjects}
    if request.method == "POST":
        if current_user.role not in staff_roles:
            abort(403)
        action = request.form.get("action", "save_grade")
        if action == "update_grade_remark":
            grade = db.session.get(Grade, request.form.get("grade_id", type=int)) or abort(404)
            if current_user.role != "teacher" or grade.teacher_id != current_user.id or grade.is_published:
                abort(403)
            remark = request.form.get("remarks", "").strip()
            if not remark:
                flash("Enter a teacher remark for this lesson.", "danger")
            else:
                grade.remarks = remark
                reviewers = User.query.filter(User.role.in_({"master_admin", "head_teacher"})).all()
                notify_users(
                    [reviewer.id for reviewer in reviewers], "Academics", "Lesson remark submitted",
                    f"{current_user.full_name} added a remark to {grade.subject} for "
                    f"{grade.student.full_name} ({grade.student.trade or 'Trade not set'}, "
                    f"{grade.student.class_name or 'Class not set'}), {grade.term or 'Term'} "
                    f"{grade.academic_year or 'Academic year not set'}: {remark}", "/grades")
                db.session.commit()
                flash("Lesson remark sent to the Master Admin and Head Teacher.", "success")
            return redirect(url_for("grades"))
        if action == "confirm_report":
            if current_user.role not in {"master_admin", "head_teacher"}:
                abort(403)
            student_id = request.form.get("student_id", type=int)
            year = request.form.get("academic_year", "").strip()
            student = db.session.get(User, student_id) or abort(404)
            valid_year = (len(year) == 9 and year[4:5] == "-" and year[:4].isdigit()
                          and year[5:].isdigit() and int(year[5:]) == int(year[:4]) + 1)
            if student.role != "student" or not valid_year:
                abort(400)
            report_grades = Grade.query.filter_by(student_id=student.id, academic_year=year).all()
            if not report_grades:
                flash("There are no saved results to confirm for this student and academic year.", "warning")
            elif any(not (grade.remarks or "").strip() for grade in report_grades):
                flash("Every subject needs teacher remarks before the report can be confirmed.", "danger")
            else:
                approval = ReportCardApproval.query.filter_by(student_id=student.id, academic_year=year).first()
                if approval is None:
                    approval = ReportCardApproval(student_id=student.id, academic_year=year, confirmed_by=current_user.id)
                    db.session.add(approval)
                else:
                    approval.confirmed_by = current_user.id
                    approval.confirmed_at = datetime.datetime.utcnow()
                for grade in report_grades:
                    grade.is_published = True
                notify_student_and_parent(
                    student, "Academics", "Report card confirmed",
                    f"Your report card for {year} has been confirmed and is ready to view.", "/grades")
                db.session.commit()
                flash(f"Report card confirmed for {student.full_name} ({year}).", "success")
            return redirect(url_for("grades"))

        if action == "bulk_save_grade":
            if current_user.role != "teacher":
                abort(403)
            class_name = request.form.get("bulk_class_name", "").strip()
            subject = request.form.get("bulk_subject", "").strip()
            academic_year = request.form.get("academic_year", "").strip()
            term = request.form.get("term", "Term 1").strip() or "Term 1"
            module_category = request.form.get("bulk_module_category", "General").strip()
            try:
                max_marks = float(request.form.get("bulk_max_marks", "100"))
            except ValueError:
                max_marks = 0
            student_ids = request.form.getlist("student_ids[]")
            marks_values = request.form.getlist("bulk_marks[]")
            remarks_values = request.form.getlist("bulk_remarks[]")
            valid_year = (len(academic_year) == 9 and academic_year[4:5] == "-"
                          and academic_year[:4].isdigit() and academic_year[5:].isdigit()
                          and int(academic_year[5:]) == int(academic_year[:4]) + 1)
            assignment_keys = {(assigned_class.casefold(), assigned_subject.casefold())
                               for assigned_class, assigned_subject in assigned_subjects}
            roster_records = User.query.filter(
                User.id.in_([int(value) for value in student_ids if value.isdigit()]),
                User.role == "student", db.func.lower(User.class_name) == class_name.casefold()
            ).all() if student_ids else []
            roster_by_id = {str(student.id): student for student in roster_records}
            grades_to_save = []
            invalid_bulk_marks = False
            for index, student_id_value in enumerate(student_ids):
                student = roster_by_id.get(student_id_value)
                if student is None:
                    invalid_bulk_marks = True
                    break
                try:
                    marks = float(marks_values[index])
                except (ValueError, IndexError):
                    invalid_bulk_marks = True
                    break
                remarks = remarks_values[index].strip() if index < len(remarks_values) else ""
                if not remarks or marks < 0 or max_marks <= 0 or marks > max_marks:
                    invalid_bulk_marks = True
                    break
                grades_to_save.append(Grade(
                    student_id=student.id, teacher_id=current_user.id, subject=subject, term=term,
                    academic_year=academic_year, module_category=module_category, marks=marks,
                    max_marks=max_marks, remarks=remarks, is_published=False))
            if (not valid_year or not class_name or (class_name.casefold(), subject.casefold()) not in assignment_keys
                    or module_category not in {"Complementary", "General", "Specific"}
                    or not student_ids or len(roster_records) != len(student_ids) or len(set(student_ids)) != len(student_ids)
                    or len(marks_values) != len(student_ids) or len(remarks_values) != len(student_ids)
                    or invalid_bulk_marks):
                flash("Choose an assigned class and lesson, then enter valid marks and teacher remarks for each listed student.", "danger")
            else:
                db.session.add_all(grades_to_save)
                reviewers = User.query.filter(User.role.in_({"master_admin", "head_teacher"})).all()
                notify_users(
                    [reviewer.id for reviewer in reviewers], "Academics", "Teacher class marks submitted",
                    f"{current_user.full_name} submitted {len(grades_to_save)} {subject} result(s) for {class_name}, "
                    f"{term} {academic_year}. These marks are ready for report review and confirmation.", "/grades")
                db.session.commit()
                flash(f"Marks for {len(grades_to_save)} students sent to the Master Admin and Head Teacher for report review.", "success")
                return redirect(url_for("grades"))
            return redirect(url_for("grades"))

        student_id = request.form.get("student_id", type=int)
        term = request.form.get("term", "Term 1").strip() or "Term 1"
        academic_year = request.form.get("academic_year", "").strip()
        field_names = ("subject", "module_code", "module_category", "module_weight", "marks", "max_marks",
                       "formative_marks", "integrated_marks", "comprehensive_marks", "reassessment_marks", "remarks")
        submitted = {name: request.form.getlist(f"{name}[]") for name in field_names}
        # Continue accepting the former one-lesson form while the multi-lesson form is in use.
        if not submitted["subject"]:
            submitted = {name: [request.form.get(name, "")] for name in field_names}
        row_count = len(submitted["subject"])
        for name in field_names:
            submitted[name] += [""] * (row_count - len(submitted[name]))

        data_rows = []
        invalid_marks = False
        for index in range(row_count):
            raw = {name: submitted[name][index].strip() for name in field_names}
            try:
                max_marks = float(raw["max_marks"] or 100)
                optional_numbers = {
                    name: float(raw[name]) if raw[name] else None
                    for name in ("module_weight", "formative_marks", "integrated_marks",
                                 "comprehensive_marks", "reassessment_marks")
                }
                if raw["marks"]:
                    marks = float(raw["marks"])
                elif any(optional_numbers[name] is not None for name in
                         ("formative_marks", "integrated_marks", "comprehensive_marks")):
                    marks = sum(optional_numbers[name] or 0 for name in
                                ("formative_marks", "integrated_marks", "comprehensive_marks"))
                else:
                    raise ValueError("Enter component marks or a total mark")
            except ValueError:
                invalid_marks = True
                marks = max_marks = 0
                optional_numbers = {name: None for name in (
                    "module_weight", "formative_marks", "integrated_marks", "comprehensive_marks", "reassessment_marks")}
            data_rows.append({**raw, "marks_value": marks, "max_marks_value": max_marks, **optional_numbers})

        valid_year = len(academic_year) == 9 and academic_year[4:5] == "-" and academic_year[:4].isdigit() and academic_year[5:].isdigit() and int(academic_year[5:]) == int(academic_year[:4]) + 1
        student = db.session.get(User, student_id) if student_id else None
        filter_trade = request.form.get("filter_trade", "")
        filter_level = request.form.get("filter_level", "")
        filter_class = request.form.get("filter_class", "")
        inferred_trade, inferred_level = CLASS_TRADE_LEVEL.get(student.class_name, (None, None)) if student else (None, None)
        filters_match = bool(student and
            (student.trade or inferred_trade or "__none__") == filter_trade and
            (student.level or inferred_level or "__none__") == filter_level and
            (student.class_name or "__none__") == filter_class)
        if not student_id or not data_rows or any(not row["subject"] for row in data_rows):
            flash("Choose a student and enter at least one lesson title.", "danger")
        elif not student or student.role != "student" or not filters_match:
            flash("Choose a student from the selected trade, level, and class.", "danger")
        elif any(not row["remarks"] for row in data_rows):
            flash("Add teacher remarks for every lesson before sending results for review.", "danger")
        elif current_user.role == "teacher" and any(((student.class_name or "").casefold(), row["subject"].casefold()) not in assigned_subject_keys for row in data_rows):
            flash("You can enter marks only for lessons assigned to you for the selected student's class.", "danger")
        elif not valid_year or any(row["module_category"] not in {"Complementary", "General", "Specific"} for row in data_rows):
            flash("Enter an academic year such as 2025-2026 and choose a valid module category.", "danger")
        elif invalid_marks or any(row["max_marks_value"] <= 0 or row["marks_value"] < 0 or
                                  row["marks_value"] > row["max_marks_value"] for row in data_rows):
            flash("Enter valid marks for every lesson. Marks must be between zero and the maximum marks.", "danger")
        elif any(value is not None and value < 0 for row in data_rows for value in
                 (row["module_weight"], row["formative_marks"], row["integrated_marks"],
                  row["comprehensive_marks"], row["reassessment_marks"])):
            flash("Module weight and assessment marks cannot be negative.", "danger")
        else:
            for row in data_rows:
                db.session.add(Grade(
                    student_id=student.id, teacher_id=current_user.id, subject=row["subject"], term=term,
                    academic_year=academic_year, module_code=row["module_code"] or None,
                    module_category=row["module_category"], module_weight=row["module_weight"],
                    marks=row["marks_value"], max_marks=row["max_marks_value"],
                    formative_marks=row["formative_marks"], integrated_marks=row["integrated_marks"],
                    comprehensive_marks=row["comprehensive_marks"], reassessment_marks=row["reassessment_marks"],
                    remarks=row["remarks"], is_published=False))
            if current_user.role == "teacher":
                reviewers = User.query.filter(User.role.in_({"master_admin", "head_teacher"})).all()
                lesson_summary = "; ".join(
                    f"{row['subject']}: {row['marks_value']:g}/{row['max_marks_value']:g} â€” {row['remarks']}"
                    for row in data_rows)
                notify_users(
                    [reviewer.id for reviewer in reviewers], "Academics", "Teacher results submitted",
                    f"{current_user.full_name} submitted {len(data_rows)} lesson result(s) for "
                    f"{student.full_name} ({student.trade or 'Trade not set'}, {student.level or 'Level not set'}, "
                    f"{student.class_name or 'Class not set'}), {term} {academic_year}. {lesson_summary}",
                    "/grades")
            db.session.commit()
            flash(f"{len(data_rows)} lesson results and teacher remarks sent to the Master Admin and Head Teacher for report review.", "success")
            return redirect(url_for("grades"))
    student_query = User.query.filter_by(role="student")
    if current_user.role == "teacher":
        assigned_classes = {class_name.casefold() for class_name, _ in assigned_subjects}
        if not assigned_classes and current_user.class_name:
            assigned_classes = {current_user.class_name.casefold()}
        student_query = student_query.filter(db.func.lower(User.class_name).in_(assigned_classes)) if assigned_classes else student_query.filter_by(id=-1)
    students = student_query.order_by(User.full_name).all() if current_user.role in {"master_admin", "head_teacher", "teacher"} else []
    grade_query = Grade.query
    if current_user.role == "teacher":
        grade_query = grade_query.filter_by(teacher_id=current_user.id)
    elif current_user.role == "student":
        grade_query = grade_query.filter_by(student_id=current_user.id)
    elif current_user.role == "parent":
        grade_query = grade_query.filter(Grade.student_id.in_([child.id for child in linked_children(current_user.id)]))
    if current_user.role in {"student", "parent", "secretary"}:
        grade_query = grade_query.filter_by(is_published=True)
    grade_list = grade_query.order_by(Grade.id.desc()).all()
    student_data = []
    for student in students:
        inferred_trade, inferred_level = CLASS_TRADE_LEVEL.get(student.class_name, (None, None))
        student_data.append({"id": student.id, "name": student.full_name,
                             "trade": student.trade or inferred_trade or "__none__",
                             "level": student.level or inferred_level or "__none__",
                             "class_name": student.class_name or "__none__"})
    student_data_json = json.dumps(student_data, ensure_ascii=False).replace("</", "<\\/")
    academic_catalog = {trade: {level: list(class_names) for level, class_names in levels.items()}
                        for trade, levels in CLASS_GROUPS.items()}
    for student in students:
        trade = student.trade or "__none__"
        level = student.level or "__none__"
        class_name = student.class_name or "__none__"
        academic_catalog.setdefault(trade, {}).setdefault(level, [])
        if class_name not in academic_catalog[trade][level]:
            academic_catalog[trade][level].append(class_name)
    academic_catalog_json = json.dumps(academic_catalog, ensure_ascii=False).replace("</", "<\\/")
    def select_options(values):
        return "".join(f"<option value='{escape(value)}'>{escape('Not specified' if value == '__none__' else value)}</option>" for value in sorted(set(values), key=lambda value: (value == "__none__", value.casefold())))
    trade_options = select_options(academic_catalog.keys())
    level_options = select_options(level for levels in academic_catalog.values() for level in levels)
    class_options_for_students = select_options(class_name for levels in academic_catalog.values()
                                                for classes in levels.values() for class_name in classes)
    student_options = "".join(f"<option value='{student.id}'>{escape(student.full_name)}</option>" for student in students)
    student_selector_html = """<div class="col-md-3"><label>Trade / Department</label><select class="form-select" name="filter_trade" id="grade-trade" required><option value="">Choose trade...</option>TRADE_OPTIONS</select></div><div class="col-md-2"><label>Level</label><select class="form-select" name="filter_level" id="grade-level" required><option value="">Choose level...</option>LEVEL_OPTIONS</select></div><div class="col-md-3"><label>Class</label><select class="form-select" name="filter_class" id="grade-class" required><option value="">Choose class...</option>CLASS_OPTIONS</select></div><div class="col-md-4"><label>Student</label><select class="form-select" name="student_id" id="grade-student" required><option value="">Choose student...</option>STUDENT_OPTIONS</select></div><script>(function(){const students=STUDENT_DATA,catalog=ACADEMIC_CATALOG;const trade=document.getElementById('grade-trade'),level=document.getElementById('grade-level'),classSelect=document.getElementById('grade-class'),student=document.getElementById('grade-student');const label=value=>value==='__none__'?'Not specified':value;function fill(select,values,placeholder){select.innerHTML='';select.add(new Option(placeholder,''));[...new Set(values)].sort((a,b)=>label(a).localeCompare(label(b))).forEach(value=>select.add(new Option(label(value),value)));select.disabled=values.length===0;}function reset(select,placeholder){select.innerHTML='';select.add(new Option(placeholder,''));select.disabled=true;}fill(trade,Object.keys(catalog),'Choose trade...');fill(level,[],'Choose trade first...');fill(classSelect,[],'Choose level first...');reset(student,'Choose class first...');trade.addEventListener('change',()=>{reset(level,'Choose trade first...');reset(classSelect,'Choose level first...');reset(student,'Choose class first...');if(!trade.value)return;fill(level,Object.keys(catalog[trade.value]||{}),'Choose level...');});level.addEventListener('change',()=>{reset(classSelect,'Choose level first...');reset(student,'Choose class first...');if(!level.value)return;fill(classSelect,(catalog[trade.value]||{})[level.value]||[],'Choose class...');});classSelect.addEventListener('change',()=>{reset(student,'Choose student...');if(!classSelect.value)return;const matches=students.filter(item=>item.trade===trade.value&&item.level===level.value&&item.class_name===classSelect.value);if(matches.length){matches.sort((a,b)=>a.name.localeCompare(b.name)).forEach(item=>student.add(new Option(item.name,item.id)));}else{student.add(new Option('No registered students in this class',''));}student.disabled=false;});})();</script>"""
    student_selector_html = (student_selector_html.replace("TRADE_OPTIONS", trade_options)
        .replace("LEVEL_OPTIONS", level_options)
        .replace("CLASS_OPTIONS", class_options_for_students)
        .replace("STUDENT_OPTIONS", student_options)
        .replace("ACADEMIC_CATALOG", academic_catalog_json)
        .replace("STUDENT_DATA", student_data_json))
    subject_options = "".join(f"<option value='{escape(subject)}' data-class='{escape(class_name)}'>{escape(subject)} Â· {escape(class_name)}</option>"
                               for class_name, subject in sorted(assigned_subjects))
    teacher_bulk_form = ""
    if current_user.role == "teacher":
        bulk_year_start = datetime.date.today().year if datetime.date.today().month >= 9 else datetime.date.today().year - 1
        bulk_default_year = f"{bulk_year_start}-{bulk_year_start + 1}"
        assigned_class_names = sorted({class_name for class_name, _ in assigned_subjects}, key=str.casefold)
        bulk_class_placeholder = "Choose assigned class..." if assigned_class_names else "No assigned classes"
        bulk_class_disabled = "" if assigned_class_names else " disabled"
        bulk_class_status = ("" if assigned_class_names else
                             "<p class='muted' role='status'>No classes or lessons are assigned to your account yet. Ask the school administrator to assign them before entering marks.</p>")
        bulk_class_options = "".join(
            f"<option value='{escape(class_name)}'>{escape(CLASS_TRADE_LEVEL.get(class_name, ('', ''))[0])} · "
            f"{escape(CLASS_TRADE_LEVEL.get(class_name, ('', ''))[1])} · {escape(class_name)}</option>"
            for class_name in assigned_class_names
        )
        bulk_lessons_json = json.dumps(
            [{"class_name": class_name, "subject": subject} for class_name, subject in sorted(assigned_subjects)],
            ensure_ascii=False,
        ).replace("</", "<\\/")
        teacher_bulk_form = """<section class='panel'><h2>Enter Marks for an Assigned Class</h2><p class='muted'>Choose one of your assigned lessons, then enter marks and remarks for the class roster in one submission.</p><form method='post' id='teacher-bulk-grade-form'><input type='hidden' name='action' value='bulk_save_grade'><div class='row g-3 mb-3'><div class='col-md-4'><label>Assigned Trade / Level / Class</label><select class='form-select' id='bulk-grade-class' name='bulk_class_name' required><option value=''>Choose assigned class...</option>BULK_CLASS_OPTIONS</select></div><div class='col-md-4'><label>Assigned Lesson</label><select class='form-select' id='bulk-grade-subject' name='bulk_subject' required disabled><option value=''>Choose class first...</option></select></div><div class='col-md-4'><label>Academic Year</label><input class='form-control' name='academic_year' pattern='[0-9]{4}-[0-9]{4}' value='BULK_YEAR' required></div><div class='col-md-3'><label>Term</label><select class='form-select' name='term'><option>Term 1</option><option>Term 2</option><option>Term 3</option></select></div><div class='col-md-3'><label>Module Category</label><select class='form-select' name='bulk_module_category'><option>Complementary</option><option selected>General</option><option>Specific</option></select></div><div class='col-md-3'><label>Maximum Marks</label><input class='form-control' type='number' name='bulk_max_marks' min='0.01' step='any' value='100' required></div></div><div class='table-responsive'><table class='table align-middle'><thead><tr><th>Student</th><th>Marks</th><th>Teacher Remarks</th></tr></thead><tbody id='bulk-grade-roster'><tr><td colspan='3' class='muted'>Choose an assigned class to load its students.</td></tr></tbody></table></div><button class='btn btn-accent' type='submit'>Send Class Marks for Review</button></form><script>(function(){const lessons=BULK_LESSONS,students=BULK_STUDENTS,classSelect=document.getElementById('bulk-grade-class'),subjectSelect=document.getElementById('bulk-grade-subject'),roster=document.getElementById('bulk-grade-roster');function loadStudents(){roster.innerHTML='';const matches=students.filter(item=>item.class_name===classSelect.value);if(!matches.length){const row=roster.insertRow();const cell=row.insertCell();cell.colSpan=3;cell.className='muted';cell.textContent='No students are registered in this assigned class yet.';return;}matches.sort((a,b)=>a.name.localeCompare(b.name)).forEach(item=>{const row=roster.insertRow(),nameCell=row.insertCell(),marksCell=row.insertCell(),remarksCell=row.insertCell(),hidden=document.createElement('input'),marks=document.createElement('input'),remarks=document.createElement('textarea');nameCell.textContent=item.name;hidden.type='hidden';hidden.name='student_ids[]';hidden.value=item.id;nameCell.appendChild(hidden);marks.type='number';marks.name='bulk_marks[]';marks.min='0';marks.step='any';marks.required=true;marks.className='form-control';marksCell.appendChild(marks);remarks.name='bulk_remarks[]';remarks.required=true;remarks.rows=1;remarks.placeholder='Teacher remark';remarks.className='form-control';remarksCell.appendChild(remarks);});}classSelect.addEventListener('change',()=>{subjectSelect.innerHTML='<option value="">Choose assigned lesson...</option>';lessons.filter(item=>item.class_name===classSelect.value).forEach(item=>subjectSelect.add(new Option(item.subject,item.subject)));subjectSelect.disabled=!classSelect.value||subjectSelect.options.length<2;loadStudents();});})();</script></section>"""
        teacher_bulk_form = (teacher_bulk_form
                             .replace("Assigned Trade / Level / Class", "Assigned Class")
                             .replace("in one submission.</p><form", "in one submission.</p>BULK_CLASS_STATUS<form")
                             .replace("name='bulk_class_name' required>", "name='bulk_class_name' requiredBULK_CLASS_DISABLED>")
                             .replace("<option value=''>Choose assigned class...</option>BULK_CLASS_OPTIONS", "<option value=''>BULK_CLASS_PLACEHOLDER</option>BULK_CLASS_OPTIONS")
                             .replace("BULK_CLASS_PLACEHOLDER", bulk_class_placeholder)
                             .replace("BULK_CLASS_DISABLED", bulk_class_disabled)
                             .replace("BULK_CLASS_STATUS", bulk_class_status)
                             .replace("BULK_CLASS_OPTIONS", bulk_class_options)
                             .replace("BULK_LESSONS", bulk_lessons_json)
                             .replace("BULK_STUDENTS", student_data_json)
                             .replace("BULK_YEAR", bulk_default_year))
    card_groups = {}
    for grade in grade_list:
        key = (grade.student_id, grade.academic_year or "Academic year not set", grade.term or "Term 1")
        card_groups.setdefault(key, []).append(grade)
    card_rows = []
    for (_, year, term), entries in card_groups.items():
        total_marks = sum(entry.marks or 0 for entry in entries)
        total_max = sum(entry.max_marks or 0 for entry in entries)
        average = (total_marks / total_max * 100) if total_max else 0
        name = escape(entries[0].student.full_name if entries[0].student else "-")
        term = escape(term)
        remarks = "<br>".join(f"<strong>{escape(entry.subject)}:</strong> {escape(entry.remarks)}" for entry in entries if entry.remarks) or "-"
        card_rows.append(f"<tr><td>{name}</td><td>{escape(year)}</td><td>{term}</td><td>{len(entries)}</td><td>{total_marks:g} / {total_max:g}</td><td>{average:.1f}%</td><td>{remarks}</td></tr>")
    report_query = Grade.query.filter_by(is_published=True)
    if current_user.role == "teacher":
        report_query = report_query.filter_by(teacher_id=current_user.id)
    elif current_user.role == "student":
        report_query = report_query.filter_by(student_id=current_user.id)
    elif current_user.role == "parent":
        report_query = report_query.filter(Grade.student_id.in_([child.id for child in linked_children(current_user.id)]))
    published_reports = report_query.order_by(Grade.academic_year.desc(), Grade.student_id, Grade.subject).all()
    report_groups = {}
    for entry in published_reports:
        report_groups.setdefault((entry.student_id, entry.academic_year), []).append(entry)
    report_card_rows = "".join(
        f"<tr><td>{escape(entries[0].student.full_name if entries[0].student else '-')}</td><td>{escape(entries[0].student.class_name if entries[0].student and entries[0].student.class_name else '-')}</td><td>{escape(year or 'Academic year not set')}</td><td class='d-flex gap-2 flex-wrap'><a class='btn btn-sm btn-outline-success' href='{url_for('report_card', student_id=student_id, year_key=year or 'unassigned')}'>View / Print Report</a><a class='btn btn-sm btn-accent' href='{url_for('export_report_card', student_id=student_id, year_key=year or 'unassigned')}'>Export Excel</a></td></tr>"
        for (student_id, year), entries in report_groups.items()
    )
    grade_rows_list = []
    for grade in grade_list:
        remarks_cell = escape(grade.remarks or "â€”")
        if current_user.role == "teacher" and grade.teacher_id == current_user.id and not grade.remarks and not grade.is_published:
            remarks_cell = (f"<form method='post' class='d-flex flex-column gap-2'><input type='hidden' name='action' value='update_grade_remark'>"
                            f"<input type='hidden' name='grade_id' value='{grade.id}'><textarea class='form-control form-control-sm' name='remarks' placeholder='Teacher remark for this lesson' required></textarea>"
                            "<button class='btn btn-sm btn-outline-success'>Send remark</button></form>")
        grade_rows_list.append(
            f"<tr><td>{escape(grade.student.full_name if grade.student else '-')}</td><td>{escape(grade.subject)}</td>"
            f"<td>{escape(grade.term or '-')}</td><td>{grade.marks:g} / {grade.max_marks or 0:g}</td><td>{remarks_cell}</td>"
            f"<td>{escape(grade.teacher.full_name if grade.teacher else '-')}</td><td>{'Published' if grade.is_published else 'Draft'}</td></tr>")
    grade_rows = "".join(grade_rows_list)
    confirmation_rows = ""
    if current_user.role in {"master_admin", "head_teacher"}:
        all_grade_rows = Grade.query.join(User, Grade.student_id == User.id).filter(User.role == "student").order_by(User.trade, User.level, User.class_name, User.full_name).all()
        confirmation_groups = {}
        for entry in all_grade_rows:
            confirmation_groups.setdefault((entry.student_id, entry.academic_year or ""), []).append(entry)
        for (sid, year), entries in confirmation_groups.items():
            student = entries[0].student
            approval = ReportCardApproval.query.filter_by(student_id=sid, academic_year=year).first() if year else None
            report_ready = bool(year) and all((entry.remarks or "").strip() for entry in entries)
            state = f"Confirmed by {escape(approval.confirmer.full_name)}" if approval and approval.confirmer else "Ready to confirm" if report_ready else "Needs academic year or teacher remarks"
            if report_ready:
                action_cell = f"<form method='post'><input type='hidden' name='action' value='confirm_report'><input type='hidden' name='student_id' value='{sid}'><input type='hidden' name='academic_year' value='{escape(year)}'><button class='btn btn-sm btn-accent'>{'Confirm again' if approval else 'Confirm report'}</button></form>"
            else:
                action_cell = "â€”"
            confirmation_rows += (f"<tr><td>{escape(student.full_name)}</td><td>{escape(student.trade or 'â€”')}</td><td>{escape(student.level or 'â€”')}</td>"
                                  f"<td>{escape(student.class_name or 'â€”')}</td><td>{escape(year or 'Not set')}</td><td>{len(entries)} subjects</td>"
                                  f"<td>{state}</td><td>{action_cell}</td></tr>")
    confirmation_body = confirmation_rows or "<tr><td colspan='8' class='muted'>No saved student results are awaiting report generation.</td></tr>"
    report_confirmation_html = ("<section class='panel'><h2>Confirm Report Generation</h2>"
        "<p class='muted'>Only the Master Admin and Head Teacher can confirm report cards. Confirmation publishes all saved subject results for that student and academic year across every trade and department.</p>"
        f"<div class='table-responsive'><table class='table align-middle'><thead><tr><th>Student</th><th>Trade / Department</th><th>Level</th><th>Class</th><th>Academic Year</th><th>Results</th><th>Status</th><th>Action</th></tr></thead><tbody>{confirmation_body}</tbody></table></div></section>"
        if current_user.role in {"master_admin", "head_teacher"} else "")
    school_year_start = datetime.date.today().year if datetime.date.today().month >= 9 else datetime.date.today().year - 1
    default_academic_year = f"{school_year_start}-{school_year_start + 1}"
    subject_control = (f"<select class='form-select' name='subject[]' required><option value=''>Choose assigned lesson...</option>{subject_options}</select>"
                       if current_user.role == "teacher" else
                       "<input class='form-control' name='subject[]' required>")
    class_ranking_html = ""
    if current_user.role in {"master_admin", "head_teacher"}:
        class_query = User.query.filter(User.role == "student", User.class_name.isnot(None), User.class_name != "")
        if current_user.role == "teacher":
            class_query = class_query.filter_by(class_name=current_user.class_name)
        class_names = sorted({name for (name,) in class_query.with_entities(User.class_name).distinct().all()})
        selected_class = request.args.get("rank_class", "")
        if current_user.role == "teacher":
            selected_class = current_user.class_name or ""
        if selected_class not in class_names:
            selected_class = ""
        selected_year = request.args.get("rank_year", default_academic_year)
        selected_term = request.args.get("rank_term", "Annual")
        if selected_term not in {"Annual", "Term 1", "Term 2", "Term 3"}:
            selected_term = "Annual"
        year_query = Grade.query.join(User, Grade.student_id == User.id).filter(
            Grade.is_published.is_(True), User.role == "student")
        if selected_class:
            year_query = year_query.filter(User.class_name == selected_class)
        if current_user.role == "teacher":
            year_query = year_query.filter(User.class_name == current_user.class_name)
        available_years = sorted({grade_year for (grade_year,) in year_query.with_entities(Grade.academic_year).distinct().all() if grade_year}, reverse=True)
        if selected_year not in available_years and selected_year != "unassigned":
            selected_year = default_academic_year if default_academic_year in available_years else (available_years[0] if available_years else default_academic_year)
        class_options = "".join(
            f"<option value='{escape(name)}' {'selected' if name == selected_class else ''}>{escape(name)}</option>"
            for name in class_names)
        year_values = list(available_years)
        if selected_year == "unassigned" or not year_values:
            year_values.append("unassigned")
        if selected_year not in year_values:
            year_values.insert(0, selected_year)
        year_options = "".join(
            f"<option value='{escape(year)}' {'selected' if year == selected_year else ''}>{escape(year if year != 'unassigned' else 'Unassigned')}</option>"
            for year in year_values)
        term_options = "".join(
            f"<option {'selected' if term == selected_term else ''}>{term}</option>"
            for term in ("Annual", "Term 1", "Term 2", "Term 3"))
        ranking_rows = ""
        if selected_class:
            ranking_query = Grade.query.join(User, Grade.student_id == User.id).filter(
                Grade.is_published.is_(True), User.role == "student", User.class_name == selected_class)
            if selected_year == "unassigned":
                ranking_query = ranking_query.filter(Grade.academic_year.is_(None))
            else:
                ranking_query = ranking_query.filter(Grade.academic_year == selected_year)
            if selected_term != "Annual":
                ranking_query = ranking_query.filter(Grade.term == selected_term)
            rank_totals = {}
            for entry in ranking_query.all():
                total = rank_totals.setdefault(entry.student_id, [entry.student, 0.0, 0.0])
                total[1] += entry.marks or 0
                total[2] += entry.max_marks or 0
            ordered_ranks = sorted(
                ((student, marks, maximum, marks / maximum * 100)
                 for student, marks, maximum in rank_totals.values() if maximum),
                key=lambda result: (-result[3], result[0].full_name.casefold()))
            previous_percent = None
            position = 0
            for index, (ranked_student, marks, maximum, percentage) in enumerate(ordered_ranks, start=1):
                if percentage != previous_percent:
                    position = index
                previous_percent = percentage
                ranking_rows += (f"<tr><td>{position}</td><td>{escape(ranked_student.full_name)}</td>"
                                 f"<td>{marks:g} / {maximum:g}</td><td>{percentage:.2f}%</td></tr>")
        ranking_content = ranking_rows or "<tr><td colspan='4' class='muted'>No published marks are available for this selection.</td></tr>"
        class_selector = (f"<input type='hidden' name='rank_class' value='{escape(selected_class)}'>"
                          if current_user.role == "teacher" else
                          f"<div class='col-md-4'><label>Class</label><select class='form-select' name='rank_class' required><option value=''>Choose class...</option>{class_options}</select></div>")
        class_ranking_html = ("<section class='panel'><h2>Class Ranking by Marks</h2>"
                              "<p class='muted'>Ranks use published results for the selected class, academic year, and term.</p>"
                              f"<form method='get' action='/grades' class='row g-3 mb-3 ranking-filter-form'>{class_selector}"
                              f"<div class='col-md-4'><label>Academic Year</label><select class='form-select' name='rank_year'>{year_options}</select></div>"
                              f"<div class='col-md-3'><label>Period</label><select class='form-select' name='rank_term'>{term_options}</select></div>"
                              "<div class='col-md-1 d-flex align-items-end'><button class='btn btn-accent'>Show</button></div></form>"
                              f"<div class='table-responsive'><table class='table'><tr><th>Rank</th><th>Student</th><th>Total Marks</th><th>Percentage</th></tr>{ranking_content}</table></div></section>")
    grade_form = """<form method="post"><input type="hidden" name="action" value="save_grade"><section class="row g-3 mb-4"><div class="col-md-6"><label>Student</label><select class="form-select" name="student_id" required><option value="">Choose student...</option>{student_options}</select></div><div class="col-md-3"><label>Academic Year</label><input class="form-control" name="academic_year" value="{default_academic_year}" pattern="[0-9]{4}-[0-9]{4}" required></div><div class="col-md-3"><label>Term</label><select class="form-select" name="term"><option>Term 1</option><option>Term 2</option><option>Term 3</option></select></div></section><h2>Lessons / Modules</h2><p class="muted">Enter raw marks for each lesson. Total marks automatically sum formative, integrated, and comprehensive assessment scores. Published results are compiled into a comprehensive report card with totals, decisions, and class position.</p><div id="lesson-rows"><div class="lesson-entry border rounded p-3 mb-3"><div class="row g-3"><div class="col-md-5"><label>Lesson / Module Title</label><input class="form-control" name="subject[]" required></div><div class="col-md-2"><label>Module Code</label><input class="form-control" name="module_code[]"></div><div class="col-md-3"><label>Module Category</label><select class="form-select" name="module_category[]"><option>Complementary</option><option>General</option><option selected>Specific</option></select></div><div class="col-md-2"><label>Module Weight</label><input type="number" min="0" step="any" class="form-control" name="module_weight[]"></div><div class="col-md-3"><label>Total Marks (auto-summed)</label><input type="number" min="0" step="any" class="form-control" name="marks[]"></div><div class="col-md-3"><label>Maximum Marks</label><input type="number" min="0.01" step="any" class="form-control" name="max_marks[]" value="100" required></div><div class="col-md-3"><label>Formative Assessment</label><input type="number" min="0" step="any" class="form-control assessment-component" name="formative_marks[]"></div><div class="col-md-3"><label>Integrated Assessment</label><input type="number" min="0" step="any" class="form-control assessment-component" name="integrated_marks[]"></div><div class="col-md-3"><label>Comprehensive Assessment</label><input type="number" min="0" step="any" class="form-control assessment-component" name="comprehensive_marks[]"></div><div class="col-md-3"><label>Re-assessment</label><input type="number" min="0" step="any" class="form-control" name="reassessment_marks[]"></div><div class="col-md-10"><label>Teacher Remarks</label><textarea class="form-control" name="remarks[]" rows="2" placeholder="Teacher remarks for this lesson" required></textarea></div><div class="col-md-2 d-flex align-items-end"><button type="button" class="btn btn-outline-danger remove-lesson">Remove lesson</button></div></div></div></div><template id="lesson-template"><div class="lesson-entry border rounded p-3 mb-3"><div class="row g-3"><div class="col-md-5"><label>Lesson / Module Title</label><input class="form-control" name="subject[]" required></div><div class="col-md-2"><label>Module Code</label><input class="form-control" name="module_code[]"></div><div class="col-md-3"><label>Module Category</label><select class="form-select" name="module_category[]"><option>Complementary</option><option>General</option><option selected>Specific</option></select></div><div class="col-md-2"><label>Module Weight</label><input type="number" min="0" step="any" class="form-control" name="module_weight[]"></div><div class="col-md-3"><label>Total Marks (auto-summed)</label><input type="number" min="0" step="any" class="form-control" name="marks[]"></div><div class="col-md-3"><label>Maximum Marks</label><input type="number" min="0.01" step="any" class="form-control" name="max_marks[]" value="100" required></div><div class="col-md-3"><label>Formative Assessment</label><input type="number" min="0" step="any" class="form-control assessment-component" name="formative_marks[]"></div><div class="col-md-3"><label>Integrated Assessment</label><input type="number" min="0" step="any" class="form-control assessment-component" name="integrated_marks[]"></div><div class="col-md-3"><label>Comprehensive Assessment</label><input type="number" min="0" step="any" class="form-control assessment-component" name="comprehensive_marks[]"></div><div class="col-md-3"><label>Re-assessment</label><input type="number" min="0" step="any" class="form-control" name="reassessment_marks[]"></div><div class="col-md-10"><label>Teacher Remarks</label><textarea class="form-control" name="remarks[]" rows="2" placeholder="Teacher remarks for this lesson" required></textarea></div><div class="col-md-2 d-flex align-items-end"><button type="button" class="btn btn-outline-danger remove-lesson">Remove lesson</button></div></div></div></template><button type="button" id="add-lesson" class="btn btn-outline-success mb-3">Add another lesson</button><p class="muted">Saved results remain drafts until the Master Admin or Head Teacher confirms report generation.</p><button class="btn btn-accent">Submit Results for Review</button></form><script>(function(){const rows=document.getElementById('lesson-rows');const template=document.getElementById('lesson-template');const classSelect=document.getElementById('grade-class');function filterAssignedLessons(){const selectedClass=classSelect?.value||'';rows.querySelectorAll('select[name="subject[]"]').forEach(select=>{[...select.options].forEach(option=>{if(option.dataset.class)option.hidden=!selectedClass||option.dataset.class!==selectedClass});if(select.selectedOptions[0]?.hidden)select.value='';});}classSelect?.addEventListener('change',filterAssignedLessons);document.getElementById('grade-student')?.addEventListener('change',filterAssignedLessons);document.getElementById('add-lesson').addEventListener('click',function(){rows.appendChild(template.content.cloneNode(true));filterAssignedLessons();});rows.addEventListener('input',function(event){if(!event.target.matches('.assessment-component'))return;const lesson=event.target.closest('.lesson-entry');const scores=[...lesson.querySelectorAll('.assessment-component')].map(input=>input.value.trim()).filter(value=>value!=='').map(Number);if(scores.length)lesson.querySelector('[name="marks[]"]').value=scores.reduce((sum,value)=>sum+value,0);});rows.addEventListener('click',function(event){if(event.target.closest('.remove-lesson')){event.target.closest('.lesson-entry').remove();}});})();</script>"""
    grade_form = (grade_form.replace('<div class="col-md-6"><label>Student</label><select class="form-select" name="student_id" required><option value="">Choose student...</option>{student_options}</select></div>', student_selector_html)
                  .replace("{default_academic_year}", default_academic_year)
                  .replace('<input class="form-control" name="subject[]" required>', subject_control)
                  if current_user.role in staff_roles else "")
    performance_rows_html = "".join(card_rows) or "<tr><td colspan='7' class='muted'>No published report cards are available.</td></tr>"
    published_report_rows_html = report_card_rows or "<tr><td colspan='4' class='muted'>No published results are available to build a report card.</td></tr>"
    performance_section = "" if current_user.role == "teacher" else (
        f"<section class='panel'><h2>Academic Performance by Term</h2><div class='table-responsive'><table class='table'><tr><th>Student</th><th>Academic Year</th><th>Term</th><th>Subjects</th><th>Total Marks</th><th>Average</th><th>Teacher Remarks</th></tr>{performance_rows_html}</table></div></section>")
    published_reports_section = "" if current_user.role == "teacher" else (
        f"<section class='panel'><h2>Published Learner Assessment Reports</h2><div class='table-responsive'><table class='table'><tr><th>Learner</th><th>Class</th><th>Academic Year</th><th>Report</th></tr>{published_report_rows_html}</table></div></section>")
    grade_page_title = "Submit Student Marks for Review" if current_user.role == "teacher" else "Academics & Report Cards"
    body = f"""<section class="panel"><h1>{escape(grade_page_title)}</h1>{teacher_bulk_form}{grade_form}</section>
{report_confirmation_html}
{class_ranking_html}
{performance_section}
{published_reports_section}
<section class="panel"><h2>{'Grade Records' if current_user.role in staff_roles else 'Published Subject Results'}</h2><div class="table-responsive"><table class="table"><tr><th>Student</th><th>Subject</th><th>Term</th><th>Marks</th><th>Teacher Remarks</th><th>Teacher</th><th>Publication</th></tr>{grade_rows or '<tr><td colspan="7" class="muted">No grade records are available.</td></tr>'}</table></div></section>"""
    return page("Grades", body, "teacher")

def get_report_card_data(student_id, year_key):
    if current_user.role == "teacher":
        abort(403)
    student = db.session.get(User, student_id)
    if not student or student.role != "student":
        abort(404)
    if current_user.role == "student" and student.id != current_user.id:
        abort(403)
    if current_user.role == "parent" and student.parent_id != current_user.id:
        abort(403)
    query = Grade.query.filter_by(student_id=student.id, is_published=True)
    academic_year = None if year_key == "unassigned" else year_key
    if academic_year is None:
        query = query.filter(Grade.academic_year.is_(None))
    else:
        query = query.filter_by(academic_year=academic_year)
    grades_for_card = query.order_by(Grade.subject, Grade.term, Grade.id).all()
    return student, academic_year, grades_for_card


def export_report_card(student_id, year_key):
    student, academic_year, grades_for_card = get_report_card_data(student_id, year_key)

    def term_index(term):
        text = (term or "").casefold()
        for index in (1, 2, 3):
            if str(index) in text or (index == 1 and "first" in text) or (index == 2 and "second" in text) or (index == 3 and "third" in text):
                return index
        return None

    modules = {}
    for grade in grades_for_card:
        index = term_index(grade.term)
        if index is None:
            continue
        key = (grade.module_code or grade.subject).strip().casefold()
        module = modules.setdefault(key, {
            "code": grade.module_code or "-", "title": grade.subject,
            "category": grade.module_category or "Specific", "weight": grade.module_weight,
            "terms": {1: [], 2: [], 3: []}, "remarks": [], "reassessment": []})
        module["terms"][index].append(grade)
        if grade.remarks and grade.remarks not in module["remarks"]:
            module["remarks"].append(grade.remarks)
        if grade.reassessment_marks is not None:
            module["reassessment"].append(grade.reassessment_marks)

    def total(records, field):
        values = [getattr(record, field) for record in records if getattr(record, field) is not None]
        return sum(values) if values else None

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Report Summary"
    summary.append(["Learner Assessment Report"])
    summary.append(["Student", student.full_name])
    summary.append(["SDMS Code", student.sdms_code or "-"])
    summary.append(["Class", student.class_name or "-"])
    summary.append(["Level", student.level or "-"])
    summary.append(["Trade", student.trade or "-"])
    summary.append(["Academic Year", academic_year or "Academic year not set"])
    peer_query = Grade.query.join(User, Grade.student_id == User.id).filter(
        Grade.is_published.is_(True), User.role == "student", User.class_name == student.class_name)
    if academic_year is None:
        peer_query = peer_query.filter(Grade.academic_year.is_(None))
    else:
        peer_query = peer_query.filter(Grade.academic_year == academic_year)
    peer_totals = {}
    for peer_grade in peer_query.all():
        peer = peer_totals.setdefault(peer_grade.student_id, [0.0, 0.0])
        peer[0] += peer_grade.marks or 0
        peer[1] += peer_grade.max_marks or 0
    target = peer_totals.get(student.id)
    class_position = "-"
    if target and target[1]:
        target_percent = target[0] / target[1]
        rank = 1 + sum(score / maximum > target_percent for score, maximum in peer_totals.values() if maximum)
        class_position = f"{rank} of {sum(bool(maximum) for _, maximum in peer_totals.values())}"
    summary.append(["Class Position", class_position])
    summary.append([])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 54

    sheet = workbook.create_sheet("Report Card")
    headers = ["Module Code", "Module Title", "Category", "Weight"]
    for term in ("Term 1", "Term 2", "Term 3"):
        headers.extend([f"{term} Formative", f"{term} Integrated", f"{term} Comprehensive", f"{term} Total"])
    headers.extend(["Annual Total", "Annual %", "Re-assessment", "Decision", "Teacher Remarks"])
    sheet.append(headers)
    category_order = {"Complementary": 0, "General": 1, "Specific": 2}
    overall_marks = overall_max = 0.0
    for module in sorted(modules.values(), key=lambda item: (category_order.get(item["category"], 9), item["title"].casefold())):
        row = [module["code"], module["title"], module["category"], module["weight"]]
        annual_marks = annual_max = 0.0
        for term in (1, 2, 3):
            records = module["terms"][term]
            for field in ("formative_marks", "integrated_marks", "comprehensive_marks"):
                row.append(total(records, field))
            marks = total(records, "marks")
            maximum = total(records, "max_marks")
            row.append(marks)
            annual_marks += marks or 0
            annual_max += maximum or 0
        percent = annual_marks / annual_max if annual_max else None
        threshold = 50 if module["category"] == "Complementary" else 70
        decision = "C" if percent is not None and percent * 100 >= threshold else "NYC" if percent is not None else "-"
        row.extend([annual_marks or None, percent, module["reassessment"][-1] if module["reassessment"] else None,
                    decision, " | ".join(module["remarks"])])
        sheet.append(row)
        overall_marks += annual_marks
        overall_max += annual_max
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in {"A": 16, "B": 34, "C": 18, "D": 12}.items():
        sheet.column_dimensions[column].width = width
    for column in range(5, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 17
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row_index, 18).number_format = "0.0%"
    summary.append(["Total Marks", overall_marks])
    summary.append(["Maximum Marks", overall_max])
    summary.append(["Overall Percentage", overall_marks / overall_max if overall_max else None])
    summary.cell(summary.max_row, 2).number_format = "0.00%"
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = secure_filename(f"report_card_{student.username}_{academic_year or 'unassigned'}.xlsx")
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def report_card(student_id, year_key):
    student, academic_year, grades_for_card = get_report_card_data(student_id, year_key)

    def term_index(term):
        text = (term or "").casefold()
        for index in (1, 2, 3):
            if str(index) in text or (index == 1 and "first" in text) or (index == 2 and "second" in text) or (index == 3 and "third" in text):
                return index
        return None

    modules = {}
    for grade in grades_for_card:
        index = term_index(grade.term)
        if index is None:
            continue
        key = (grade.module_code or grade.subject).strip().casefold()
        module = modules.setdefault(key, {
            "code": grade.module_code or "-", "title": grade.subject,
            "category": grade.module_category or "Specific", "weight": grade.module_weight,
            "terms": {1: [], 2: [], 3: []}, "remarks": {1: [], 2: [], 3: []},
            "reassessment": [],
        })
        module["terms"][index].append(grade)
        if grade.remarks:
            module["remarks"][index].append(grade.remarks)
        if grade.reassessment_marks is not None:
            module["reassessment"].append(grade.reassessment_marks)

    def sum_values(records, field):
        values = [getattr(record, field) for record in records if getattr(record, field) is not None]
        return sum(values) if values else None

    def number(value):
        return "â€”" if value is None else f"{value:g}"

    def score_cell(records, field):
        value = sum_values(records, field)
        return number(value)

    category_order = {"Complementary": 0, "General": 1, "Specific": 2}
    grouped_modules = {}
    for module in modules.values():
        category = module["category"] if module["category"] in category_order else "Specific"
        grouped_modules.setdefault(category, []).append(module)
    for category_modules in grouped_modules.values():
        category_modules.sort(key=lambda item: (item["title"].casefold(), item["code"].casefold()))

    rows = []
    term_totals = {1: [0.0, 0.0], 2: [0.0, 0.0], 3: [0.0, 0.0]}
    term_components = {i: {"formative_marks": [], "integrated_marks": [], "comprehensive_marks": []} for i in (1, 2, 3)}
    annual_total = 0.0
    annual_maximum = 0.0
    for category in sorted(grouped_modules, key=lambda value: category_order.get(value, 9)):
        rows.append(f"<tr class='category-row'><th colspan='19'>{escape(category)} Modules</th></tr>")
        for module in grouped_modules[category]:
            module_annual = 0.0
            module_maximum = 0.0
            cells = []
            for index in (1, 2, 3):
                records = module["terms"][index]
                for field in ("formative_marks", "integrated_marks", "comprehensive_marks"):
                    value = sum_values(records, field)
                    if value is not None:
                        term_components[index][field].append(value)
                    cells.append(f"<td>{number(value)}</td>")
                score = sum_values(records, "marks")
                maximum = sum_values(records, "max_marks")
                if score is not None:
                    module_annual += score
                    term_totals[index][0] += score
                if maximum is not None:
                    module_maximum += maximum
                    term_totals[index][1] += maximum
                cells.append(f"<td class='total-mark'>{number(score)}</td>")
            annual_total += module_annual
            annual_maximum += module_maximum
            annual_percent = (module_annual / module_maximum * 100) if module_maximum else None
            module_reassessment = module["reassessment"][-1] if module["reassessment"] else None
            threshold = 50 if category == "Complementary" else 70
            decision = "â€”" if annual_percent is None else "C" if annual_percent >= threshold else "NYC"
            rows.append(
                f"<tr><td>{escape(module['code'])}</td><td>{escape(module['title'])}</td><td>{number(module['weight'])}</td>"
                f"{''.join(cells)}<td>{number(module_annual)}</td><td>{f'{annual_percent:.1f}%' if annual_percent is not None else 'â€”'}</td>"
                f"<td>{number(module_reassessment)}</td><td>{decision}</td></tr>"
            )

    def pct(score, maximum):
        return f"{(score / maximum * 100):.2f}%" if maximum else "â€”"

    total_cells = []
    percentage_cells = []
    for index in (1, 2, 3):
        for field in ("formative_marks", "integrated_marks", "comprehensive_marks"):
            values = term_components[index][field]
            total_cells.append(f"<td>{number(sum(values) if values else None)}</td>")
        score, maximum = term_totals[index]
        total_cells.append(f"<td class='total-mark'>{number(score) if maximum else 'â€”'}</td>")
        percentage_cells.append(f"<td colspan='4'>{pct(score, maximum)}</td>")
    overall_percent = pct(annual_total, annual_maximum)
    rows.append(f"<tr class='summary-row'><th colspan='3'>TOTAL</th>{''.join(total_cells)}<td>{number(annual_total) if annual_maximum else 'â€”'}</td><td>{overall_percent}</td><td>â€”</td><td>â€”</td></tr>")
    rows.append(f"<tr class='summary-row'><th colspan='3'>PERCENTAGE</th>{''.join(percentage_cells)}<td colspan='4'>{overall_percent}</td></tr>")

    peer_query = Grade.query.join(User, Grade.student_id == User.id).filter(
        Grade.is_published.is_(True), User.role == "student", User.class_name == student.class_name
    )
    if academic_year is None:
        peer_query = peer_query.filter(Grade.academic_year.is_(None))
    else:
        peer_query = peer_query.filter(Grade.academic_year == academic_year)
    peer_grades = peer_query.all()
    peer_term_values = {}
    peer_year_values = {}
    for peer_grade in peer_grades:
        peer_term = term_index(peer_grade.term)
        if peer_term is None:
            continue
        per_student_term = peer_term_values.setdefault((peer_grade.student_id, peer_term), [0.0, 0.0])
        per_student_term[0] += peer_grade.marks or 0
        per_student_term[1] += peer_grade.max_marks or 0
        per_student_year = peer_year_values.setdefault(peer_grade.student_id, [0.0, 0.0])
        per_student_year[0] += peer_grade.marks or 0
        per_student_year[1] += peer_grade.max_marks or 0
    positions = []
    for index in (1, 2, 3):
        target = peer_term_values.get((student.id, index))
        if not target or not target[1]:
            positions.append("â€”")
        else:
            target_pct = target[0] / target[1]
            ranked = [v[0] / v[1] for (sid, term), v in peer_term_values.items() if term == index and v[1]]
            positions.append(f"{1 + sum(value > target_pct for value in ranked)} Out of {len(ranked)}")
    target_year = peer_year_values.get(student.id)
    annual_position = "â€”"
    if target_year and target_year[1]:
        target_pct = target_year[0] / target_year[1]
        ranked = [value[0] / value[1] for value in peer_year_values.values() if value[1]]
        annual_position = f"{1 + sum(value > target_pct for value in ranked)} Out of {len(ranked)}"
    position_cells = "".join(f"<td colspan='4'>{value}</td>" for value in positions)
    rows.append(f"<tr><th colspan='3'>POSITION</th>{position_cells}<td colspan='4'>{annual_position}</td></tr>")

    term_comments = []
    for index in (1, 2, 3):
        comments = []
        for module in modules.values():
            for comment in module["remarks"][index]:
                if comment not in comments:
                    comments.append(comment)
        term_comments.append("<br>".join(escape(comment) for comment in comments) or "â€”")
    annual_comments = "<br>".join(escape(grade.remarks) for grade in grades_for_card if grade.remarks) or "â€”"
    comments_cells = "".join(f"<td colspan='4'>{comment}</td>" for comment in term_comments)
    rows.append(f"<tr><th colspan='3'>Class Trainer's Comments & Signature</th>{comments_cells}<td colspan='4'>{annual_comments}</td></tr>")
    rows_html = "".join(rows)
    school_name = escape(get_setting("school_name", "Lycée Saint Alexandre Sauli de Muhura"))
    school_district = escape(get_setting("school_district", "Gatsibo"))
    school_email = escape(get_setting("school_email", "lycemuhura@gmail.com"))
    school_phone = escape(get_setting("school_phone", "0788862998"))
    year_label = academic_year or "Academic year not set"
    class_label = (student.class_name or "").strip()
    program_trade = re.sub(r"^(?:L|S)\d+\s+", "", class_label, flags=re.IGNORECASE)
    program_trade = re.sub(r"\s+[AB]$", "", program_trade, flags=re.IGNORECASE) or student.trade or "â€”"
    level_value = student.level or CLASS_TRADE_LEVEL.get(class_label, (None, None))[1] or "â€”"
    level_digits = re.search(r"\d+", level_value)
    rtqf_level = level_digits.group(0) if level_digits else "â€”"
    qualification = ("TVET CERTIFICATE III" if rtqf_level == "4" else
                     "TVET CERTIFICATE II" if rtqf_level == "3" else
                     "TVET CERTIFICATE IV" if rtqf_level == "5" else "TVET QUALIFICATION")
    behavior_row = "<tr class='behavior-row'><td colspan='2'>Behaviour</td><td>40</td>" + "".join("<td>â€”</td>" for _ in range(16)) + "</tr>"
    report_html = f"""<style>
.report-card{{background:#fff;color:#111;padding:12px;border:1px solid #555;font-family:Arial,sans-serif;font-size:9px;max-width:1200px;margin:auto}}
.report-card h1,.report-card h2,.report-card p{{margin:0}}.report-head{{display:grid;grid-template-columns:1fr 110px 1fr;gap:12px;align-items:center;margin-bottom:12px;font-size:10px}}
.report-head img{{width:105px;height:105px;object-fit:contain;margin:auto}}.report-title{{text-align:center;font-weight:700;border:1px solid #222;padding:5px;margin-bottom:0}}
.report-meta,.report-table{{width:100%;border-collapse:collapse}}.report-meta{{margin-bottom:5px}}.report-meta td{{border:1px solid #777;padding:4px;font-weight:600}}
.report-table{{table-layout:fixed;font-size:7px}}.report-table th,.report-table td{{border:1px solid #555;padding:3px;text-align:center;overflow-wrap:anywhere}}
.report-table th:nth-child(1){{width:6%}}.report-table th:nth-child(2){{width:17%}}.report-table th:nth-child(3){{width:4%}}
.report-table thead th{{background:#f4f4f4}}.report-table .vertical-heading{{height:150px;min-width:18px;writing-mode:vertical-rl;transform:rotate(180deg);white-space:nowrap;vertical-align:middle;padding:5px 2px}}
.report-table .category-row th{{background:#e8e8e8;font-size:8px}}.report-table .behavior-row,.report-table .total-mark,.report-table .summary-row{{font-weight:700}}
.report-signoff{{display:grid;grid-template-columns:1fr 1fr;gap:12px;margin-top:8px;font-size:8px}}.deliberation,.school-manager{{border:1px solid #555;padding:7px;min-height:100px}}.decision-options{{display:grid;grid-template-columns:1fr auto;gap:4px 8px}}.signature{{margin-top:14px;border-top:1px solid #555;padding-top:4px;min-height:20px}}.report-legend{{display:grid;grid-template-columns:1fr 1fr;gap:12px;margin-top:6px;font-size:8px}}
.report-controls{{max-width:1200px;margin:0 auto 12px;display:flex;gap:8px}}
@media print{{@page{{size:A4 portrait;margin:6mm}}body{{background:#fff!important}}.sidebar,.topbar,.flash,footer,.report-controls{{display:none!important}}.main-shell{{margin:0!important}}.content{{padding:0!important;max-width:none!important}}.report-card{{max-width:none;border:0;padding:0;font-size:8px}}.report-head{{font-size:8px;grid-template-columns:1fr 90px 1fr}}.report-head img{{width:88px;height:88px}}.report-meta td{{padding:3px}}.report-table{{font-size:6px}}.report-table th,.report-table td{{padding:2px}}.report-table .vertical-heading{{height:138px}}}}
</style>
<div class="report-controls"><button class="btn btn-accent" onclick="window.print()">Print / Save as PDF</button><a class="btn btn-outline-success" href="{url_for('export_report_card', student_id=student.id, year_key=year_key)}">Export Excel</a><a class="btn btn-outline-success" href="/grades">Back to Results</a></div>
<article class="report-card"><header class="report-head"><div><strong>REPUBLIC OF RWANDA</strong><br>MINISTRY OF EDUCATION<br>DISTRICT: {school_district.upper()}<br><strong>{school_name.upper()}</strong><br>Email: {school_email}<br>Tel: {school_phone}</div><img src="{get_school_logo_static_url()}" alt="School logo" onerror="this.style.display='none'"><div><strong>ACADEMIC YEAR</strong>ã€€{escape(year_label)}<br><strong>CLASS</strong>ã€€{escape(student.class_name or 'â€”')}<br><strong>LEARNER'S NAME</strong>ã€€{escape(student.full_name)}<br><strong>SDMS CODE</strong>ã€€{escape(student.sdms_code or 'â€”')}</div></header>
<table class="report-meta"><tr><td colspan="4" style="text-align:center">LEARNER'S ASSESSMENT REPORT</td></tr><tr><td>Sector</td><td>{escape(get_setting('school_sector','ICT AND MULTIMEDIA'))}</td><td>Qualification Title</td><td>{escape(qualification + ' IN ' + program_trade.upper())}</td></tr><tr><td>Trade</td><td>{escape(program_trade.upper())}</td><td>RTQF Level</td><td>{escape(rtqf_level)}</td></tr></table>
<table class="report-table"><thead><tr><th rowspan="2">Module<br>Code</th><th rowspan="2">Module Title</th><th rowspan="2">Module Weight</th><th colspan="4">1ST TERM</th><th colspan="4">2ND TERM</th><th colspan="4">3RD TERM</th><th colspan="4">ANNUAL</th></tr><tr><th class="vertical-heading">Formative Assessment</th><th class="vertical-heading">Integrated Assessment</th><th class="vertical-heading">Comprehensive Assessment</th><th class="vertical-heading">Total Marks</th><th class="vertical-heading">Formative Assessment</th><th class="vertical-heading">Integrated Assessment</th><th class="vertical-heading">Comprehensive Assessment</th><th class="vertical-heading">Total Marks</th><th class="vertical-heading">Formative Assessment</th><th class="vertical-heading">Integrated Assessment</th><th class="vertical-heading">Comprehensive Assessment</th><th class="vertical-heading">Total Marks</th><th class="vertical-heading">Annual Total</th><th class="vertical-heading">Annual %</th><th class="vertical-heading">Re-assessment</th><th class="vertical-heading">Decision</th></tr></thead><tbody>{behavior_row}{rows_html or '<tr><td colspan="19">No published subject results for this academic year.</td></tr>'}</tbody></table>
<div class="report-signoff"><div class="deliberation"><strong>Deliberation</strong><div class="decision-options"><span>Promoted at 1st sitting</span><span>â–¡</span><span>Promoted after re-assessment</span><span>â–¡</span><span>Advised to repeat</span><span>â–¡</span><span>Dismissed</span><span>â–¡</span></div><div class="signature">Date: ______________________________</div></div><div class="school-manager"><strong>School Manager</strong><div class="signature">Name & Signature: ______________________________</div><div class="signature">Official stamp</div></div></div>
<div class="report-legend"><div>Module Weight = Module's learning hours Ã— credit.<br>Passing line: 50% for Mathematics, Sciences and Complementary Modules.<br>Passing line: 70% for Core Modules (Specific and General Modules).</div><div><strong>Decision</strong><br>C: Competent<br>NYC: Not Yet Competent</div></div></article>"""
    return page("Learner Assessment Report", report_html, "student")

def register_routes(app):
    app.add_url_rule('/academic-management', endpoint='academic_management', view_func=permission_for('academic_management')(academic_management), methods=['GET', 'POST'])
    app.add_url_rule('/timetable', endpoint='timetable', view_func=permission_for('timetable')(timetable), methods=['GET', 'POST'])
    app.add_url_rule('/day-schedule', endpoint='day_schedule', view_func=permission_for('day_schedule')(day_schedule), methods=['GET', 'POST'])
    app.add_url_rule('/grades', endpoint='grades', view_func=permission_for('grades')(grades), methods=['GET', 'POST'])
    app.add_url_rule('/grades/report/<int:student_id>/<year_key>', endpoint='report_card', view_func=permission_for('report_card')(report_card), methods=['GET'])
    app.add_url_rule('/grades/report/<int:student_id>/<year_key>/export.xlsx', endpoint='export_report_card', view_func=permission_for('export_report_card')(export_report_card), methods=['GET'])
