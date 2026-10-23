import calendar
import datetime
import io
import json
import secrets
import textwrap

from flask import abort, current_app, flash, jsonify, redirect, request, Response, send_file, session, url_for
from flask_login import current_user
from markupsafe import escape

from app.extensions import db
from app.models import (
    Meeting, MeetingActionItem, MeetingAgendaItem, MeetingAuditEvent,
    MeetingAttendance, MeetingMinutes, MeetingMinutesRevision,
    MeetingNotificationLog, MeetingParticipant, MessageConversation,
    Setting, StudentGuardian, User,
)
from app.services.meeting_service import (
    ACTION_STATUSES, APPROVER_ROLES, ATTENDANCE_STATUSES, CLASS_SUPPORT_ROLES,
    INVITATION_STATUSES,
    MEETING_ROLES, MEETING_STATUSES, OFFICE_ROLES, can_manage_meeting, can_view_meeting,
    escalate_overdue_actions, meeting_reminders, staff_classes,
    guardian_ids_for_student, record_meeting_audit, verified_student_ids,
    visible_meetings,
)
from app.services.meeting_workflow import (
    add_action_comment, add_agenda_item, add_meeting_participant,
    create_action_item, create_meeting,
    decide_meeting, decide_minutes, record_attendance, remove_agenda_item,
    move_agenda_item,
    respond_to_invitation, revoke_guardian_link, save_minutes,
    submit_meeting_for_approval, submit_minutes, update_action,
    update_meeting, verify_guardian_link, revoke_meeting_participant,
)
from app.services.room_booking import (
    check_in_student_for_current_meetings, current_room_reservations,
    extend_room_reservation, release_room_reservation, release_unchecked_meeting_rooms,
    reserve_room, suggested_room_names,
)
from app.services.microsoft_calendar import sync_meeting_to_microsoft
from app.utils.decorators import permission_for
from app.utils.helpers import guardian_ids_for_student, now_string
from app.utils.rendering import page


def _csrf_token():
    token = session.get("meeting_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["meeting_csrf_token"] = token
    return token


def _csrf_valid():
    provided = request.headers.get("X-CSRF-Token") if request.is_json else request.form.get("csrf_token")
    expected = session.get("meeting_csrf_token", "")
    return bool(provided and expected and secrets.compare_digest(provided, expected))


def _csrf_field():
    return f"<input type='hidden' name='csrf_token' value='{escape(_csrf_token())}'>"


def _get_meeting(meeting_id):
    meeting = db.session.get(Meeting, meeting_id)
    if not meeting:
        abort(404)
    return meeting


def _meeting_thread(meeting):
    return MessageConversation.query.filter_by(meeting_id=meeting.id).first()


def _format_time(value):
    try:
        return datetime.time.fromisoformat(value).strftime("%H:%M")
    except (ValueError, TypeError):
        return value or ""


def _students_for_user(user):
    query = User.query.filter_by(role="student", is_active_account=True)
    if user.role in OFFICE_ROLES:
        return query.order_by(User.class_name, User.full_name).all()
    classes = staff_classes(user)
    return query.filter(db.func.lower(User.class_name).in_(classes)).order_by(
        User.class_name, User.full_name
    ).all() if classes else []


def _meeting_form(user):
    students = _students_for_user(user)
    classes = sorted({student.class_name for student in students if student.class_name}, key=str.casefold)
    student_options = "".join(
        f"<option value='{student.id}'>{escape(student.full_name)} · {escape(student.class_name or 'No class')}</option>"
        for student in students
    )
    class_options = "".join(f"<option value='{escape(name)}'>{escape(name)}</option>" for name in classes)
    parents = User.query.filter_by(role="parent", is_active_account=True).order_by(User.full_name).all()
    staff = User.query.filter(
        User.role.in_(OFFICE_ROLES | {"teacher", "discipline", "patron", "matron"}),
        User.is_active_account.is_(True),
    ).order_by(User.full_name).all()
    if user.role == "teacher" or user.role in CLASS_SUPPORT_ROLES:
        allowed_classes = staff_classes(user)
        parent_ids = {
            guardian_id for student in students
            for guardian_id in guardian_ids_for_student(student.id)
        }
        parents = [parent for parent in parents if parent.id in parent_ids]
        staff = [candidate for candidate in staff if candidate.role in OFFICE_ROLES or
                 candidate.id == user.id or (candidate.class_name and
                 candidate.class_name.casefold() in allowed_classes)]
    participant_options = "".join(
        f"<option value='{person.id}' data-role='{escape(person.role)}'>{escape(person.full_name)} · {escape(person.role.replace('_', ' ').title())}</option>"
        for person in parents + staff if person.id != user.id
    )
    approver_options = "".join(
        f"<option value='{person.id}'>{escape(person.full_name)} · {escape(person.role.replace('_', ' ').title())}</option>"
        for person in staff if person.role in APPROVER_ROLES
    )
    type_options = "".join(
        f"<option value='{value}'>{label}</option>"
        for value, label in (
            ("general", "General parents' meeting"), ("class", "Class meeting"),
            ("parent_teacher", "Parent–teacher appointment"), ("disciplinary", "Disciplinary meeting"),
            ("emergency", "Emergency meeting"), ("committee", "Authorized committee"),
            ("panel_discussion", "Panel discussion"),
        )
    )
    moderator_options = "".join(
        f"<option value='{person.id}'>{escape(person.full_name)} Â· {escape(person.role.replace('_', ' ').title())}</option>"
        for person in staff if person.role in OFFICE_ROLES | {"teacher"}
    )
    panelist_options = "".join(
        f"<option value='{person.id}'>{escape(person.full_name)} Â· {escape(person.role.replace('_', ' ').title())}</option>"
        for person in staff if person.role in OFFICE_ROLES | {"teacher", "discipline", "patron", "matron"}
    )
    publish_default = " checked"
    return f"""<details class='panel meeting-create'><summary><strong>Create a meeting</strong><span class='muted ms-2'>Schedule, invite participants, and prepare an agenda</span></summary>
<form method='post' class='row g-3 mt-2'><input type='hidden' name='action' value='create'>{_csrf_field()}
<div class='col-md-6'><label class='form-label'>Title</label><input class='form-control' name='title' maxlength='160' required></div>
<div class='col-md-3'><label class='form-label'>Meeting type</label><select class='form-select' name='meeting_type' data-meeting-type-select>{type_options}</select></div>
<div class='col-12 panel-discussion-settings' data-panel-settings hidden><div class='panel panel-format-guide'><h3>Panel discussion format</h3><p class='muted'>A live, unscripted exchange with three to five staff panelists and a separate moderator. Use contrasting perspectives and open prompts instead of separate slide presentations.</p><div class='row g-2'><div class='col-md-6'><label class='form-label'>Panelists (select 3–5)</label><select class='form-select' name='panelist_ids' multiple size='5'>{panelist_options}</select></div><div class='col-md-3'><label class='form-label'>Moderator</label><select class='form-select' name='panel_moderator_id'><option value=''>Choose moderator</option>{moderator_options}</select></div><div class='col-md-3'><label class='form-label'>Audience Q&amp;A</label><select class='form-select' name='qa_duration_minutes'>{''.join(f"<option value='{minutes}'{' selected' if minutes == 15 else ''}>{minutes} minutes</option>" for minutes in range(15, 21))}</select></div></div><p class='small muted mb-0 mt-2'>Total runtime must be 30–90 minutes. The agenda reserves the selected Q&amp;A time at the end.</p></div></div>
<div class='col-12' data-panel-perspectives hidden><label class='form-label'>Perspectives to bring into the discussion</label><textarea class='form-control' name='panel_perspectives' rows='2' maxlength='1000' placeholder='For example: classroom teaching, student wellbeing, and family perspectives'></textarea></div>
<div class='col-md-3'><label class='form-label'>Student (optional)</label><select class='form-select' name='student_id'><option value=''>No specific student</option>{student_options}</select></div>
<div class='col-md-4'><label class='form-label'>Class (optional)</label><select class='form-select' name='class_name'><option value=''>No class scope</option>{class_options}</select></div>
<div class='col-md-4'><label class='form-label'>Date</label><input class='form-control' type='date' name='scheduled_date' value='{datetime.date.today().isoformat()}' required></div>
<div class='col-md-2'><label class='form-label'>Starts</label><input class='form-control' type='time' name='start_time' required></div>
<div class='col-md-2'><label class='form-label'>Ends</label><input class='form-control' type='time' name='end_time' required></div>
<div class='col-md-4'><label class='form-label'>Room / location</label><input class='form-control' name='location' maxlength='160'></div>
<div class='col-md-4'><label class='form-label'>Invite audience</label><select class='form-select' name='audience_mode'><option value='selected'>Selected participants</option><option value='guardians'>Verified guardians of selected student</option><option value='class_parents'>Verified guardians in this class</option>{'<option value="all_parents">All active parents / guardians</option>' if user.role in OFFICE_ROLES else ''}<option value='committee'>Authorized committee participants</option></select></div>
<div class='col-md-4'><label class='form-label'>Participants</label><select class='form-select' name='participant_ids' multiple size='4'>{participant_options}</select><small class='muted'>Select parent/guardian and staff participants for the selected audience.</small></div>
<div class='col-md-4' data-standard-chairperson><label class='form-label'>Chairperson</label><select class='form-select' name='chairperson_id'><option value=''>Not assigned</option>{''.join(f"<option value='{person.id}'>{escape(person.full_name)}</option>" for person in staff)}</select></div>
<div class='col-md-4'><label class='form-label'>Meeting secretary</label><select class='form-select' name='secretary_id'><option value=''>Not assigned</option>{''.join(f"<option value='{person.id}'>{escape(person.full_name)}</option>" for person in staff)}</select></div>
<div class='col-md-4'><label class='form-label'>Minutes approver</label><select class='form-select' name='minutes_approver_id'><option value=''>Choose designated approver</option>{approver_options}</select></div>
<div class='col-md-6'><label class='form-label'>Description</label><textarea class='form-control' name='description' rows='2' maxlength='10000'></textarea></div>
<div class='col-md-6'><label class='form-label'>Agenda items (one per line)</label><textarea class='form-control' name='agenda_items' rows='2'></textarea></div>
<div class='col-md-4'><label class='form-label'>Confidentiality</label><select class='form-select' name='confidentiality'><option value='participants'>Invited participants</option><option value='staff_only'>Staff only</option></select></div>
<div class='col-md-8 d-flex align-items-end'><label><input type='checkbox' name='publish' value='yes'{publish_default}> Publish now or send for approval</label></div>
<div class='col-12'><button class='btn btn-accent'>Save meeting</button></div></form><script>(() => {{ const form=document.currentScript.previousElementSibling; const type=form.elements.meeting_type; const panel=form.querySelector('[data-panel-settings]'); const perspectives=form.querySelector('[data-panel-perspectives]'); const chair=form.querySelector('[data-standard-chairperson]'); const update=() => {{ const active=type.value==='panel_discussion'; panel.hidden=!active; perspectives.hidden=!active; chair.hidden=active; }}; type.addEventListener('change',update); update(); }})();</script></details>"""


def _meeting_card(meeting, viewer):
    participants = meeting.participants
    if viewer.role == "parent":
        own = next((row for row in participants if row.user_id == viewer.id), None)
        accepted = int(bool(own and own.invitation_status == "accepted"))
        declined = 0
        pending = int(bool(own and own.invitation_status in {"invited", "reschedule_requested"}))
    else:
        accepted = sum(row.invitation_status == "accepted" for row in participants)
        declined = sum(row.invitation_status == "declined" for row in participants)
        pending = sum(row.invitation_status in {"invited", "pending_approval", "reschedule_requested"} for row in participants)
    context = meeting.class_name or (meeting.student.class_name if meeting.student else "School-wide")
    if viewer.role == "parent" and meeting.student:
        context = meeting.student.class_name or "Student meeting"
    if viewer.role == "parent":
        response_text = (
            "Your invitation accepted" if accepted else
            "Your response is pending" if pending else "Your response recorded"
        )
    else:
        response_text = f"{accepted} accepted<br>{pending} awaiting response"
        if viewer.role in OFFICE_ROLES:
            response_text += f"<br>{declined} declined"
    room_release_note = (
        f"<p class='text-danger small mb-1'>{escape(meeting.room_release_reason or 'Room released because no attendee checked in.')}</p>"
        if meeting.room_released_at else ""
    )
    return (
        f"<article class='panel meeting-tile'><div class='d-flex justify-content-between align-items-start gap-3 flex-wrap'>"
        f"<div><span class='pill'>{escape(meeting.status)}</span>{room_release_note}<h3 class='mt-2 mb-1'><a href='{escape(url_for('meeting_detail', meeting_id=meeting.id))}'>{escape(meeting.title)}</a></h3>"
        f"<p class='muted mb-0'>{escape(context)} · {escape(meeting.meeting_type.replace('_', ' ').title())} · {escape(meeting.scheduled_date)} · {escape(_format_time(meeting.start_time))}–{escape(_format_time(meeting.end_time))}</p></div>"
        f"<div class='text-end small'>{response_text}</div></div></article>"
    )


def _calendar_html(meetings, month_value):
    try:
        year, month = (int(value) for value in month_value.split("-", 1))
        if month < 1 or month > 12:
            raise ValueError
    except (ValueError, TypeError):
        today = datetime.date.today()
        year, month = today.year, today.month
    grouped = {}
    for meeting in meetings:
        try:
            scheduled = datetime.date.fromisoformat(meeting.scheduled_date)
        except ValueError:
            continue
        if scheduled.year == year and scheduled.month == month:
            grouped.setdefault(scheduled.day, []).append(meeting)
    day_headers = "".join(f"<th scope='col'>{day}</th>" for day in ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"))
    weeks = []
    for week in calendar.Calendar(firstweekday=0).monthdayscalendar(year, month):
        cells = []
        for day in week:
            events = grouped.get(day, []) if day else []
            event_html = "".join(
                f"<a class='calendar-meeting' href='{escape(url_for('meeting_detail', meeting_id=item.id))}'>{escape(_format_time(item.start_time))} · {escape(item.title)}</a>"
                for item in events
            )
            cells.append(f"<td><strong>{day or ''}</strong>{event_html}</td>")
        weeks.append("<tr>" + "".join(cells) + "</tr>")
    return f"<div class='table-responsive'><table class='table calendar-table'><thead><tr>{day_headers}</tr></thead><tbody>{''.join(weeks)}</tbody></table></div>"


def _live_schedule_payload(user):
    now = datetime.datetime.now()
    horizon = now + datetime.timedelta(days=7)
    current, upcoming, known_locations = [], [], set()
    for meeting in visible_meetings(user):
        if meeting.status not in {"Scheduled", "Rescheduled"} or meeting.room_released_at:
            continue
        location = (meeting.location or "").strip()
        try:
            start = datetime.datetime.combine(
                datetime.date.fromisoformat(meeting.scheduled_date),
                datetime.time.fromisoformat(meeting.start_time),
            )
            end = datetime.datetime.combine(
                datetime.date.fromisoformat(meeting.scheduled_date),
                datetime.time.fromisoformat(meeting.end_time),
            )
        except (TypeError, ValueError):
            continue
        if start > horizon or end <= now:
            continue
        if location:
            known_locations.add(location)
        participants = [
            row for row in meeting.participants
            if row.invitation_status not in {"declined", "revoked"}
        ]
        event = {
            "id": meeting.id,
            "title": meeting.title,
            "location": location or "Location to be confirmed",
            "date": meeting.scheduled_date,
            "start": meeting.start_time,
            "end": meeting.end_time,
            "present": sum(bool(row.attendance and row.attendance.status == "Present") for row in participants),
            "expected": len(participants),
        }
        if start <= now < end:
            current.append(event)
        elif start > now:
            upcoming.append(event)
    occupied = {event["location"].casefold() for event in current if event["location"] != "Location to be confirmed"}
    try:
        configured_limits = current_app.config.get("ROOM_CAPACITY_LIMITS", "{}")
        configured_limits = json.loads(configured_limits) if isinstance(configured_limits, str) else configured_limits
        capacity_limits = {
            str(name).strip().casefold(): int(limit)
            for name, limit in configured_limits.items()
            if str(name).strip() and int(limit) > 0
        }
    except (AttributeError, TypeError, ValueError, json.JSONDecodeError):
        capacity_limits = {}
    occupancy_by_room = {}
    for event in current:
        if event["location"] != "Location to be confirmed":
            key = event["location"].casefold()
            occupancy_by_room[key] = occupancy_by_room.get(key, 0) + event["present"]
    rooms = [
        {
            "name": name,
            "occupied": name.casefold() in occupied,
            "present": occupancy_by_room.get(name.casefold(), 0),
            "capacity": capacity_limits.get(name.casefold()),
            "over_capacity": (
                name.casefold() in capacity_limits
                and occupancy_by_room.get(name.casefold(), 0) > capacity_limits[name.casefold()]
            ),
        }
        for name in sorted(known_locations, key=str.casefold)
    ]
    for event in current:
        capacity = capacity_limits.get(event["location"].casefold())
        event["capacity"] = capacity
        event["over_capacity"] = capacity is not None and event["present"] > capacity
    key = lambda event: (event["date"], event["start"], event["title"].casefold())
    return {
        "updated_at": now.strftime("%H:%M:%S"),
        "rooms": rooms,
        "current": sorted(current, key=key),
        "upcoming": sorted(upcoming, key=key)[:12],
    }


def meeting_live_schedule():
    release_unchecked_meeting_rooms()
    payload = _live_schedule_payload(current_user)
    active_reservations = current_room_reservations()
    rooms_by_name = {room["name"].casefold(): room for room in payload["rooms"]}
    for booking in active_reservations:
        name = booking.room_name.strip()
        key = name.casefold()
        room = rooms_by_name.setdefault(key, {"name": name, "occupied": False})
        room["occupied"] = True
    now = datetime.datetime.now()
    today_meetings = Meeting.query.filter(
        Meeting.scheduled_date == now.date().isoformat(),
        Meeting.status.in_({"Scheduled", "Rescheduled"}),
        Meeting.location.isnot(None),
    ).all()
    for meeting in today_meetings:
        if meeting.room_released_at:
            continue
        try:
            start = datetime.time.fromisoformat(meeting.start_time)
            end = datetime.time.fromisoformat(meeting.end_time)
        except (TypeError, ValueError):
            continue
        if start <= now.time() < end and meeting.location.strip():
            name = meeting.location.strip()
            room = rooms_by_name.setdefault(name.casefold(), {"name": name, "occupied": False, "present": 0, "capacity": None, "over_capacity": False})
            room["occupied"] = True
    payload["rooms"] = sorted(rooms_by_name.values(), key=lambda room: room["name"].casefold())
    return jsonify(payload)


def _room_booking_json(booking):
    return {
        "id": booking.id,
        "room": booking.room_name,
        "purpose": booking.purpose,
        "reserved_by": booking.reserved_by.full_name,
        "starts_at": booking.starts_at.strftime("%Y-%m-%d %H:%M"),
        "ends_at": booking.ends_at.strftime("%Y-%m-%d %H:%M"),
        "ends_time": booking.ends_at.strftime("%H:%M"),
        "can_manage": booking.reserved_by_id == current_user.id or current_user.role in OFFICE_ROLES,
    }


def room_booking_api():
    if current_user.role == "parent":
        abort(403)
    if request.method == "GET":
        return jsonify({
            "rooms": suggested_room_names(),
            "bookings": [_room_booking_json(booking) for booking in current_room_reservations()],
        })
    if not _csrf_valid():
        abort(400)
    values = request.get_json(silent=True) or request.form
    action = (values.get("action") or "reserve").strip()
    try:
        booking_id = int(values.get("booking_id") or 0)
    except (TypeError, ValueError):
        booking_id = 0
    try:
        if action == "reserve":
            booking = reserve_room(
                current_user, values.get("room"), values.get("purpose", ""),
                values.get("minutes", 60),
            )
            message = f"{booking.room_name} reserved until {booking.ends_at.strftime('%H:%M')}."
        elif action == "extend":
            booking = extend_room_reservation(
                current_user, booking_id, values.get("minutes", 30),
            )
            message = f"{booking.room_name} booking extended until {booking.ends_at.strftime('%H:%M')}."
        elif action == "release":
            booking = release_room_reservation(current_user, booking_id)
            message = f"{booking.room_name} released."
        else:
            return jsonify({"error": "Choose reserve, extend, or release."}), 400
    except PermissionError as error:
        return jsonify({"error": str(error)}), 403
    except ValueError as error:
        return jsonify({"error": str(error)}), 409
    return jsonify({"message": message, "booking": _room_booking_json(booking)})


def room_booking_qr():
    if current_user.role == "parent":
        abort(403)
    room_name = request.args.get("room", "").strip()
    if not room_name or len(room_name) > 160:
        abort(400)
    from qrcode import QRCode
    from qrcode.image.svg import SvgPathImage

    public_base = current_app.config.get("PUBLIC_BASE_URL") or request.url_root.rstrip("/")
    booking_url = f"{public_base}{url_for('meetings', room=room_name, view='list')}"
    qr = QRCode(box_size=8, border=4)
    qr.add_data(booking_url)
    qr.make(fit=True)
    image = qr.make_image(image_factory=SvgPathImage)
    buffer = io.BytesIO()
    image.save(buffer)
    return Response(buffer.getvalue(), mimetype="image/svg+xml", headers={"Cache-Control": "private, max-age=300"})


def _room_booking_panel(initial_room=""):
    return f"""<section class='panel room-booking-panel' data-room-booking data-endpoint='{url_for('room_booking_api')}' data-qr-endpoint='{url_for('room_booking_qr')}' data-csrf='{escape(_csrf_token())}'>
<div class='live-schedule-heading'><div><span class='pill'>TOUCHSCREEN BOOKING</span><h2>Quick room booking</h2></div><p class='muted mb-0'>Reserve a room now, add time, or release it when finished.</p></div>
<form data-room-booking-form class='quick-booking-form'><div><label class='form-label' for='quick-room-name'>Room</label><input class='form-control touch-control' id='quick-room-name' name='room' list='quick-room-options' maxlength='160' value='{escape(initial_room)}' required placeholder='Select or enter a room'><datalist id='quick-room-options'></datalist><button class='btn btn-outline-secondary mt-2 touch-button' type='button' data-show-room-qr>Show booking QR</button></div>
<div><label class='form-label' for='quick-room-purpose'>Purpose</label><input class='form-control touch-control' id='quick-room-purpose' name='purpose' maxlength='160' placeholder='Optional'></div>
<div><label class='form-label' for='quick-room-duration'>Duration</label><select class='form-select touch-control' id='quick-room-duration' name='minutes'><option value='30'>30 minutes</option><option value='60' selected>1 hour</option><option value='90'>1 hour 30 minutes</option><option value='120'>2 hours</option><option value='180'>3 hours</option><option value='240'>4 hours</option></select></div>
<button class='btn btn-accent touch-button' type='submit'>Reserve now</button></form>
<div data-room-qr-container class='room-qr-container' hidden><img data-room-qr alt='QR code to open room booking on a mobile device'><p class='muted'>Scan with a phone to open this room in the meeting panel.</p></div>
<div data-room-booking-message class='booking-message' role='status' aria-live='polite'></div><div data-room-bookings><p class='muted'>Loading active room bookings…</p></div></section>
<style>.room-qr-container{{display:flex;align-items:center;gap:14px;margin-top:12px}}.room-qr-container[hidden]{{display:none}}.room-qr-container img{{width:150px;height:150px;background:#fff;padding:8px;border-radius:8px}}@media(max-width:800px){{.room-qr-container{{flex-wrap:wrap}}}}</style>
<script>
(() => {{
 const panel = document.querySelector('[data-room-booking]'); if (!panel) return;
 const form = panel.querySelector('[data-room-booking-form]');
 const bookingsNode = panel.querySelector('[data-room-bookings]');
 const roomsNode = panel.querySelector('#quick-room-options');
 const messageNode = panel.querySelector('[data-room-booking-message]');
 const qrButton = panel.querySelector('[data-show-room-qr]');
 const qrContainer = panel.querySelector('[data-room-qr-container]');
 const qrImage = panel.querySelector('[data-room-qr]');
 const say = (text, error = false) => {{ messageNode.textContent = text; messageNode.className = 'booking-message' + (error ? ' error' : ''); }};
 const load = async () => {{
   try {{
     const response = await fetch(panel.dataset.endpoint, {{headers: {{'Accept': 'application/json'}}}});
     const data = await response.json(); if (!response.ok) throw new Error(data.error || 'Could not load room bookings.');
     roomsNode.replaceChildren(...data.rooms.map((name) => {{ const option = document.createElement('option'); option.value = name; return option; }}));
     if (!data.bookings.length) {{ bookingsNode.replaceChildren(Object.assign(document.createElement('p'), {{className: 'muted', textContent: 'No rooms are currently reserved.'}})); return; }}
     bookingsNode.replaceChildren(...data.bookings.map((booking) => {{
       const card = document.createElement('article'); card.className = 'quick-booking-card';
       const info = document.createElement('div');
       const title = document.createElement('strong'); title.textContent = booking.room; info.append(title);
       const detail = document.createElement('p'); detail.className = 'muted'; detail.textContent = booking.purpose + ' · until ' + booking.ends_time + ' · ' + booking.reserved_by; info.append(detail);
       card.append(info);
       if (booking.can_manage) {{
         const controls = document.createElement('div'); controls.className = 'quick-booking-controls';
         const extend = document.createElement('button'); extend.type = 'button'; extend.className = 'btn btn-outline-success touch-button'; extend.textContent = 'Extend 30 min';
         extend.addEventListener('click', () => update('extend', booking.id));
         const release = document.createElement('button'); release.type = 'button'; release.className = 'btn btn-outline-danger touch-button'; release.textContent = 'Release room';
         release.addEventListener('click', () => update('release', booking.id));
         controls.append(extend, release); card.append(controls);
       }}
       return card;
     }}));
   }} catch (error) {{ say(error.message, true); }}
 }};
 const update = async (action, bookingId) => {{
   const payload = action === 'reserve'
     ? {{action, room: form.elements.room.value, purpose: form.elements.purpose.value, minutes: form.elements.minutes.value}}
     : {{action, booking_id: bookingId, minutes: 30}};
   try {{
     const response = await fetch(panel.dataset.endpoint, {{method: 'POST', headers: {{'Content-Type': 'application/json', 'X-CSRF-Token': panel.dataset.csrf}}, body: JSON.stringify(payload)}});
     const data = await response.json(); if (!response.ok) throw new Error(data.error || 'Room booking could not be updated.');
     say(data.message); if (action === 'reserve') form.reset(); await load();
   }} catch (error) {{ say(error.message, true); }}
 }};
 qrButton.addEventListener('click', () => {{
   const room = form.elements.room.value.trim();
   if (!room) {{ say('Select or enter a room first.', true); return; }}
   qrImage.src = panel.dataset.qrEndpoint + '?room=' + encodeURIComponent(room);
   qrContainer.hidden = false;
 }});
 form.addEventListener('submit', (event) => {{ event.preventDefault(); update('reserve'); }});
 load(); window.setInterval(load, 60000);
}})();
</script>"""


def _live_schedule_panel():
    return f"""<section class='panel live-schedule' data-live-schedule data-endpoint='{url_for('meeting_live_schedule')}'>
<div class='live-schedule-heading'><div><span class='pill'>LIVE SCHEDULE</span><h2>Room activity</h2></div><span class='muted'>Updated <time data-live-updated>--:--:--</time> · refreshes every minute</span></div>
<h3>Room status</h3><div class='room-indicators' data-live-rooms><p class='muted'>Loading room status…</p></div>
<div class='live-schedule-columns'><div><h3>Happening now</h3><div data-live-current><p class='muted'>Loading current meetings…</p></div></div>
<div><h3>Upcoming reservations · 7 days</h3><div data-live-upcoming><p class='muted'>Loading upcoming meetings…</p></div></div></div></section>
<script>
(() => {{
  const panel = document.querySelector('[data-live-schedule]');
  if (!panel) return;
  const roomsNode = panel.querySelector('[data-live-rooms]');
  const currentNode = panel.querySelector('[data-live-current]');
  const upcomingNode = panel.querySelector('[data-live-upcoming]');
  const updatedNode = panel.querySelector('[data-live-updated]');
  const message = (text) => {{ const p = document.createElement('p'); p.className = 'muted'; p.textContent = text; return p; }};
  const eventNode = (event, active) => {{
    const item = document.createElement('article'); item.className = active ? 'live-event live-event-active' : 'live-event';
    const title = document.createElement('a'); title.href = '/meetings/' + encodeURIComponent(event.id); title.className = 'live-event-title';
    title.textContent = event.title; item.append(title);
    const details = document.createElement('p'); details.className = 'muted';
    details.textContent = event.date + ' · ' + event.start + '–' + event.end + ' · ' + event.location; item.append(details);
    const occupancy = document.createElement('span'); occupancy.className = 'live-occupancy';
    occupancy.textContent = event.present + ' present · ' + event.expected + ' invited' + (event.capacity ? ' · capacity ' + event.capacity : ''); item.append(occupancy);
    if (event.over_capacity) {{ const alert = document.createElement('p'); alert.className = 'room-capacity-alert'; alert.setAttribute('role', 'alert'); alert.textContent = 'Capacity exceeded: ' + event.present + ' present, limit ' + event.capacity + '.'; item.append(alert); }}
    return item;
  }};
  const refresh = async () => {{
    try {{
      const response = await fetch(panel.dataset.endpoint, {{headers: {{'Accept': 'application/json'}}}});
      if (!response.ok) throw new Error('Schedule refresh failed');
      const data = await response.json();
      roomsNode.replaceChildren(...(data.rooms.length ? data.rooms.map((room) => {{
        const badge = document.createElement('span'); badge.className = 'room-indicator ' + (room.occupied ? 'occupied' : 'available');
        const light = document.createElement('i'); light.className = 'status-light'; light.setAttribute('aria-hidden', 'true');
        const label = document.createElement('span'); label.textContent = room.name + ' · ' + (room.occupied ? 'Occupied' : 'Available') + (room.capacity ? ' · ' + room.present + '/' + room.capacity : '');
        if (room.over_capacity) {{ badge.classList.add('over-capacity'); badge.setAttribute('role', 'alert'); label.textContent += ' · CAPACITY EXCEEDED'; }}
        badge.append(light, label); return badge;
      }}) : [message('No rooms are listed on upcoming meetings.') ]));
      currentNode.replaceChildren(...(data.current.length ? data.current.map((event) => eventNode(event, true)) : [message('No meetings are in progress.') ]));
      upcomingNode.replaceChildren(...(data.upcoming.length ? data.upcoming.map((event) => eventNode(event, false)) : [message('No upcoming reservations in the next seven days.') ]));
      updatedNode.textContent = data.updated_at;
    }} catch (_) {{
      if (!updatedNode.textContent || updatedNode.textContent === '--:--:--') {{
        roomsNode.replaceChildren(message('Live room status is temporarily unavailable.'));
        currentNode.replaceChildren(message('Please refresh the page to try again.'));
        upcomingNode.replaceChildren(message('Please refresh the page to try again.'));
      }}
    }}
  }};
  refresh(); window.setInterval(refresh, 60000);
}})();
</script>"""


def meetings():
    if request.method == "POST":
        if not _csrf_valid():
            abort(400)
        action = request.form.get("action", "")
        if action == "create":
            try:
                meeting = create_meeting(current_user, request.form)
                if meeting.status in {"Scheduled", "Rescheduled"}:
                    sync_meeting_to_microsoft(meeting)
                flash("Meeting saved and invitations sent." if meeting.status == "Scheduled" else
                      "Meeting saved. It will be visible to invitees after approval." if meeting.status == "Pending approval" else
                      "Meeting saved as a draft.", "success")
                return redirect(url_for("meeting_detail", meeting_id=meeting.id))
            except PermissionError:
                db.session.rollback()
                abort(403)
            except ValueError as error:
                db.session.rollback()
                flash(str(error), "danger")
        else:
            abort(400)

    records = visible_meetings(current_user)
    query = request.args.get("q", "").strip().casefold()
    status_filter = request.args.get("status", "").strip()
    type_filter = request.args.get("type", "").strip()
    view = request.args.get("view", "list")
    month_value = request.args.get("month", datetime.date.today().strftime("%Y-%m"))
    if query:
        records = [meeting for meeting in records if query in meeting.title.casefold() or
                   query in (meeting.class_name or "").casefold() or
                   query in (meeting.location or "").casefold()]
    if status_filter in MEETING_STATUSES:
        records = [meeting for meeting in records if meeting.status == status_filter]
    if type_filter:
        records = [meeting for meeting in records if meeting.meeting_type == type_filter]
    upcoming = sum(meeting.status in {"Scheduled", "Rescheduled"} and meeting.scheduled_date >= datetime.date.today().isoformat() for meeting in records)
    if current_user.role == "parent":
        awaiting = sum(any(
            row.user_id == current_user.id
            and row.invitation_status in {"invited", "reschedule_requested"}
            for row in meeting.participants
        ) for meeting in records)
    else:
        awaiting = sum(meeting.status == "Pending approval" or any(
            row.invitation_status in {"invited", "reschedule_requested"} for row in meeting.participants
        ) for meeting in records)
    open_actions = sum(
        action.status not in {"Completed", "Reviewed"}
        for meeting in records for action in meeting.actions
        if current_user.role != "parent" or action.is_parent_visible
    )
    summary = f"""<section class='panel meeting-summary'><div><span class='pill'>MEETING WORKSPACE</span><h1>Parents & School <span class='accent'>Meetings</span></h1>
<p class='muted'>Manage invitations, attendance, approved minutes, and follow-up actions in one secure workspace.</p></div>
<div class='meeting-stats'><div><strong>{upcoming}</strong><span>Upcoming</span></div><div><strong>{awaiting}</strong><span>Pending approval / response</span></div><div><strong>{open_actions}</strong><span>Open actions</span></div></div></section>"""
    create_form = "" if current_user.role == "parent" else _meeting_form(current_user)
    calendar_filter = ""
    if view == "calendar":
        calendar_filter = f"""<form method='get' class='row g-2 mt-2'>
<input type='hidden' name='view' value='calendar'><input type='hidden' name='q' value='{escape(query)}'>
<div class='col-auto'><label class='form-label'>Calendar month</label><input class='form-control' type='month' name='month' value='{escape(month_value)}'></div>
<div class='col-auto d-flex align-items-end'><button class='btn btn-outline-success'>Go</button></div></form>"""
    filters = f"""<section class='panel'><form method='get' class='row g-2 align-items-end'>
<div class='col-md-4'><label class='form-label'>Search meetings</label><input class='form-control' name='q' value='{escape(request.args.get('q', ''))}' placeholder='Title, class, or room'></div>
<div class='col-md-3'><label class='form-label'>Status</label><select class='form-select' name='status'><option value=''>All statuses</option>{''.join(f"<option value='{escape(item)}{' selected' if status_filter == item else ''}'>{escape(item)}</option>" for item in sorted(MEETING_STATUSES))}</select></div>
<div class='col-md-3'><label class='form-label'>Type</label><select class='form-select' name='type'><option value=''>All types</option>{''.join(f"<option value='{item}'{' selected' if type_filter == item else ''}>{escape(item.replace('_', ' ').title())}</option>" for item in ('general','class','parent_teacher','disciplinary','emergency','committee','panel_discussion'))}</select></div>
<div class='col-md-2'><button class='btn btn-outline-success w-100'>Filter</button></div></form>
<div class='d-flex gap-2 mt-3'><a class='btn btn-sm btn-outline-secondary' href='{escape(url_for('meetings', view='list', q=query, status=status_filter, type=type_filter))}'>List view</a><a class='btn btn-sm btn-outline-secondary' href='{escape(url_for('meetings', view='calendar', month=month_value, q=query, status=status_filter, type=type_filter))}'>Calendar view</a></div>
{calendar_filter}
</section>"""
    cards = _calendar_html(records, month_value) if view == "calendar" else "".join(_meeting_card(meeting, current_user) for meeting in records)
    if not cards:
        cards = "<section class='panel muted'>No meetings match this view.</section>"
    parent_request = (
        "<section class='panel'><h2>Parent appointment requests</h2><p>Request a meeting for a linked child from the <a href='/appointments'>Parent–School Meeting Requests</a> panel.</p></section>"
        if current_user.role == "parent" else ""
    )
    admin_links = ""
    if current_user.role in OFFICE_ROLES:
        admin_links = "<section class='panel d-flex gap-2 flex-wrap'><a class='btn btn-sm btn-outline-success' href='/meetings/guardians'>Verify parent / guardian links</a><a class='btn btn-sm btn-outline-success' href='/meetings/reports'>Meeting reports</a>"
        if current_user.role in {"master_admin", "head_teacher"}:
            admin_links += "<a class='btn btn-sm btn-outline-secondary' href='/meetings/settings'>Meeting settings</a>"
        if current_user.role == "master_admin":
            admin_links += "<a class='btn btn-sm btn-outline-secondary' href='/meetings/audit'>Audit trail</a>"
        admin_links += "</section>"
    booking_panel = "" if current_user.role == "parent" else _room_booking_panel(request.args.get("room", "").strip())
    body = summary + admin_links + create_form + _live_schedule_panel() + booking_panel + filters + parent_request + f"<section class='meeting-list'>{cards}</section>" + _meeting_styles()
    return page("Meeting Panel", body, "parent" if current_user.role == "parent" else "master_admin")


def _meeting_styles():
    return """<style>
.meeting-summary{display:flex;justify-content:space-between;align-items:center;gap:24px;flex-wrap:wrap}.meeting-summary h1{margin:.5rem 0}.meeting-stats{display:flex;gap:10px;flex-wrap:wrap}.meeting-stats div{min-width:130px;border:1px solid var(--border);border-radius:12px;padding:12px 16px;display:grid}.meeting-stats strong{font-size:1.5rem;color:var(--green-700)}.meeting-stats span{font-size:.8rem;color:var(--muted)}.meeting-create summary{cursor:pointer}.meeting-tile{margin-bottom:12px}.meeting-tile a{text-decoration:none;color:inherit}.meeting-tile a:hover{color:var(--green-700)}.meeting-detail-grid{display:grid;grid-template-columns:minmax(0,1fr) minmax(300px,.7fr);gap:18px}.meeting-section{margin-bottom:16px}.calendar-table td{vertical-align:top;height:116px;min-width:100px}.calendar-meeting{display:block;font-size:.75rem;border-radius:7px;padding:4px 6px;margin:4px 0;background:rgba(8,155,112,.09);color:var(--text);overflow-wrap:anywhere}.meeting-status{white-space:nowrap}.live-schedule h2,.room-booking-panel h2{margin:.35rem 0 .8rem}.live-schedule h3{font-size:.9rem;font-weight:700;margin:1rem 0 .55rem}.live-schedule-heading{display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap}.live-schedule-columns{display:grid;grid-template-columns:1fr 1fr;gap:18px}.room-indicators{display:flex;gap:9px;flex-wrap:wrap}.room-indicator{display:inline-flex;align-items:center;gap:8px;border:1px solid var(--border);border-radius:999px;padding:7px 12px;font-size:.85rem;font-weight:700}.room-indicator.available{color:#147a43;background:#eaf8ef;border-color:#9bd8b0}.room-indicator.occupied{color:#a1262e;background:#fff0f0;border-color:#e5a3a6}.room-indicator.over-capacity{color:#7b1f00;background:#ffe4d6;border-color:#ed9675}.room-capacity-alert{color:#7b1f00;background:#ffe4d6;border-radius:6px;padding:6px 8px;font-weight:700}.status-light{width:12px;height:12px;border-radius:50%;background:#20a653;box-shadow:0 0 8px #20a653}.occupied .status-light{background:#dc3545;box-shadow:0 0 8px #dc3545}.live-event{border-left:3px solid var(--green-700);padding:8px 10px;margin:7px 0;background:var(--surface);border-radius:7px}.live-event-active{border-left-color:#dc3545}.live-event-title{font-weight:700;color:var(--text)}.live-event p{margin:.25rem 0}.live-occupancy{font-size:.78rem;color:var(--muted)}.quick-booking-form{display:grid;grid-template-columns:1.2fr 1.2fr .8fr auto;align-items:end;gap:12px;margin-top:18px}.touch-control{min-height:50px;font-size:1rem}.touch-button{min-height:50px;min-width:130px;padding:10px 16px;font-weight:700}.quick-booking-card{display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap;border:1px solid var(--border);border-radius:12px;padding:12px 14px;margin-top:10px}.quick-booking-card p{margin:.3rem 0 0}.quick-booking-controls{display:flex;gap:8px;flex-wrap:wrap}.booking-message{min-height:1.5rem;margin-top:10px;font-weight:700;color:var(--green-700)}.booking-message.error{color:#a1262e}@media(max-width:800px){.meeting-detail-grid,.live-schedule-columns,.quick-booking-form{grid-template-columns:1fr}.calendar-table td{min-width:80px;height:90px}}
</style>"""


def meeting_detail(meeting_id):
    meeting = _get_meeting(meeting_id)
    if not can_view_meeting(current_user, meeting):
        abort(403)
    if request.method == "POST":
        if not _csrf_valid():
            abort(400)
        action = request.form.get("action", "")
        try:
            if action == "update":
                update_meeting(current_user, meeting, request.form)
                flash("Meeting details updated.", "success")
            elif action == "submit_meeting":
                submit_meeting_for_approval(current_user, meeting)
                flash("Meeting sent for approval.", "success")
            elif action in {"approve_meeting", "reject_meeting"}:
                decide_meeting(current_user, meeting, "approved" if action == "approve_meeting" else "rejected",
                               request.form.get("note", ""))
                flash("Meeting approval saved.", "success")
            elif action == "respond":
                respond_to_invitation(current_user, meeting, request.form.get("response", ""),
                                      request.form.get("comment", ""))
                flash("Your meeting response was sent.", "success")
            elif action == "participant_add":
                add_meeting_participant(current_user, meeting, request.form.get("user_id", type=int))
                flash("Meeting invitation sent.", "success")
            elif action == "participant_revoke":
                revoke_meeting_participant(
                    current_user, meeting, request.form.get("participant_id", type=int),
                )
                flash("Meeting invitation withdrawn.", "success")
            elif action == "attendance":
                record_attendance(
                    current_user, meeting, request.form.get("participant_id", type=int),
                    request.form.get("status", ""), request.form.get("representative_name", ""),
                    request.form.get("representative_relationship", ""),
                )
                flash("Attendance record saved.", "success")
            elif action == "agenda_add":
                add_agenda_item(current_user, meeting, request.form.get("title", ""), request.form.get("details", ""))
                flash("Agenda item added.", "success")
            elif action == "agenda_remove":
                remove_agenda_item(current_user, meeting, request.form.get("agenda_item_id", type=int))
                flash("Agenda item removed.", "success")
            elif action == "agenda_move":
                move_agenda_item(current_user, meeting, request.form.get("agenda_item_id", type=int),
                                 request.form.get("direction", ""))
                flash("Agenda order updated.", "success")
            elif action == "minutes_save":
                save_minutes(current_user, meeting, request.form.get("content", ""))
                flash("Minutes saved as a draft.", "success")
            elif action == "minutes_submit":
                submit_minutes(current_user, meeting)
                flash("Minutes submitted for approval.", "success")
            elif action in {"minutes_approve", "minutes_reject"}:
                decide_minutes(current_user, meeting, "approved" if action == "minutes_approve" else "rejected",
                               request.form.get("note", ""))
                flash("Minutes decision saved.", "success")
            elif action == "action_create":
                create_action_item(current_user, meeting, request.form)
                flash("Follow-up action assigned.", "success")
            elif action == "action_update":
                action_item = db.session.get(MeetingActionItem, request.form.get("action_id", type=int))
                if not action_item or action_item.meeting_id != meeting.id:
                    abort(404)
                update_action(current_user, action_item, request.form)
                flash("Action progress updated.", "success")
            elif action == "action_comment":
                action_item = db.session.get(MeetingActionItem, request.form.get("action_id", type=int))
                if not action_item or action_item.meeting_id != meeting.id:
                    abort(404)
                add_action_comment(current_user, action_item, request.form.get("body", ""),
                                   request.form.get("is_parent_visible") == "yes")
                flash("Action comment added.", "success")
            else:
                abort(400)
            if action in {"update", "approve_meeting", "reject_meeting"}:
                sync_meeting_to_microsoft(meeting)
        except PermissionError:
            db.session.rollback()
            abort(403)
        except ValueError as error:
            db.session.rollback()
            flash(str(error), "danger")
        return redirect(url_for("meeting_detail", meeting_id=meeting.id))

    body = _render_meeting_detail(meeting, current_user)
    return page(meeting.title, body + _meeting_styles(), "parent" if current_user.role == "parent" else "master_admin")


def _render_meeting_detail(meeting, viewer):
    can_manage = can_manage_meeting(viewer, meeting)
    parent_view = viewer.role == "parent"
    participant = next((row for row in meeting.participants if row.user_id == viewer.id), None)
    thread = _meeting_thread(meeting)
    thread_link = (
        f"<a class='btn btn-sm btn-outline-success' href='{escape(url_for('messages', conversation=thread.id))}'>Open secure meeting conversation</a>"
        if thread and participant and participant.invitation_status != "revoked" else ""
    )
    student_text = ""
    if meeting.student:
        if not parent_view or meeting.student.id in verified_student_ids(viewer.id):
            student_text = f" · Student: {escape(meeting.student.full_name)}"
    description = f"<p>{escape(meeting.description).replace(chr(10), '<br>')}</p>" if meeting.description else ""
    panel_details = ""
    if meeting.meeting_type == "panel_discussion":
        panelists = [row.user.full_name for row in meeting.participants
                     if row.participant_role == "panelist" and row.user]
        moderator = meeting.chairperson.full_name if meeting.chairperson else "Not assigned"
        panel_details = (
            "<section class='panel meeting-section'><h2>Panel discussion</h2>"
            f"<p><strong>Moderator:</strong> {escape(moderator)}</p>"
            f"<p><strong>Panelists:</strong> {escape(', '.join(panelists) or 'Not assigned')}</p>"
            f"<p><strong>Audience Q&amp;A:</strong> {int(meeting.qa_duration_minutes or 15)} minutes</p>"
            + (f"<p><strong>Perspectives:</strong> {escape(meeting.panel_perspectives)}</p>" if meeting.panel_perspectives else "")
            + "<p class='muted mb-0'>Keep the exchange conversational and unscripted. The moderator balances airtime, guides transitions, and invites different perspectives.</p></section>"
        )
    calendar_status = ""
    if meeting.microsoft_sync_status == "synced":
        calendar_status = "<p class='muted'>Synced to Microsoft 365 calendar.</p>"
        if meeting.microsoft_join_url:
            calendar_status += f"<p><a class='btn btn-sm btn-outline-success' href='{escape(meeting.microsoft_join_url)}' target='_blank' rel='noopener'>Join Teams meeting</a></p>"
    elif meeting.microsoft_sync_status == "error":
        calendar_status = "<p class='text-danger'>Microsoft calendar sync failed. Check the meeting sync settings.</p>"
    overview = f"""<section class='panel'><div class='d-flex justify-content-between align-items-start gap-3 flex-wrap'>
<div><span class='pill'>{escape(meeting.status)} · {escape(meeting.approval_status.replace('_', ' ').title())}</span><h1 class='mt-2'>{escape(meeting.title)}</h1>
<p class='muted'>{escape(meeting.meeting_type.replace('_', ' ').title())} · {escape(meeting.scheduled_date)} · {escape(_format_time(meeting.start_time))}–{escape(_format_time(meeting.end_time))} · {escape(meeting.location or 'Location to be confirmed')}{student_text}</p>{description}</div>{thread_link}</div></section>"""
    approval_feedback = ""
    if can_manage and meeting.approval_status == "rejected":
        event = MeetingAuditEvent.query.filter_by(
            meeting_id=meeting.id, event_type="meeting_rejected",
        ).order_by(MeetingAuditEvent.id.desc()).first()
        try:
            note = json.loads(event.details).get("note", "") if event else ""
        except (TypeError, ValueError):
            note = ""
        if note:
            approval_feedback = f"<section class='panel'><h2>Approval feedback</h2><p>{escape(note)}</p></section>"

    response_form = ""
    if parent_view and participant and participant.invitation_status in {"invited", "accepted", "declined", "reschedule_requested"} and meeting.status not in {"Cancelled", "Completed"}:
        response_form = f"""<section class='panel meeting-section'><h2>Your invitation</h2><p>Current response: <strong>{escape(participant.invitation_status.replace('_', ' ').title())}</strong></p>
<form method='post' class='row g-2'><input type='hidden' name='action' value='respond'>{_csrf_field()}
<div class='col-md-4'><label class='form-label'>Response</label><select class='form-select' name='response'><option value='accepted'>Accept invitation</option><option value='declined'>Decline invitation</option><option value='reschedule_requested'>Request rescheduling</option></select></div>
<div class='col-md-6'><label class='form-label'>Comment (optional)</label><input class='form-control' name='comment' maxlength='1000'></div><div class='col-md-2 d-flex align-items-end'><button class='btn btn-accent'>Send response</button></div></form></section>"""

    participant_rows = []
    active_participant_ids = {row.user_id for row in meeting.participants if row.invitation_status != "revoked"}
    for row in meeting.participants:
        if parent_view and row.user.role == "parent" and row.user_id != viewer.id:
            continue
        attendance = row.attendance
        attendance_text = attendance.status if attendance else "Not recorded"
        representative = (
            f" · Representative: {escape(attendance.representative_name)} ({escape(attendance.representative_relationship)})"
            if attendance and attendance.representative_name else ""
        )
        response = escape(row.invitation_status.replace("_", " ").title()) if not parent_view or row.user_id == viewer.id else "Invited"
        if can_manage:
            options = "".join(
                f"<option value='{escape(status)}{' selected' if attendance and attendance.status == status else ''}>{escape(status)}</option>"
                for status in sorted(ATTENDANCE_STATUSES)
            )
            attendance_control = f"""<form method='post' class='row g-1 align-items-end'><input type='hidden' name='action' value='attendance'>{_csrf_field()}<input type='hidden' name='participant_id' value='{row.id}'>
<div class='col-md-4'><label class='form-label'>Attendance</label><select class='form-select form-select-sm' name='status'>{options}</select></div>
<div class='col-md-3'><label class='form-label'>Representative</label><input class='form-control form-control-sm' name='representative_name' value='{escape(attendance.representative_name if attendance else '')}' maxlength='160'></div>
<div class='col-md-3'><label class='form-label'>Relationship</label><input class='form-control form-control-sm' name='representative_relationship' value='{escape(attendance.representative_relationship if attendance else '')}' maxlength='80'></div>
<div class='col-md-2'><button class='btn btn-sm btn-outline-success'>Save</button></div></form>"""
        else:
            attendance_control = f"{escape(attendance_text)}{representative}"
        management_control = ""
        if can_manage and row.user_id != meeting.organizer_id and row.invitation_status != "revoked":
            management_control = f"<form method='post' onsubmit=\"return confirm('Withdraw this meeting invitation?')\"><input type='hidden' name='action' value='participant_revoke'>{_csrf_field()}<input type='hidden' name='participant_id' value='{row.id}'><button class='btn btn-sm btn-outline-danger'>Withdraw</button></form>"
        participant_rows.append(
            f"<tr><td>{escape(row.user.full_name)}</td><td>{escape(row.participant_role.replace('_', ' ').title())}</td>"
            f"<td>{response}</td><td>{attendance_control}</td><td>{management_control}</td></tr>"
        )
    participant_add_form = ""
    if can_manage and meeting.status not in {"Completed", "Cancelled"}:
        candidate_ids = set()
        if viewer.role in OFFICE_ROLES and meeting.confidentiality != "staff_only":
            candidate_ids.update(user.id for user in User.query.filter(
                User.role == "parent", User.is_active_account.is_(True),
            ).all() if verified_student_ids(user.id))
        if meeting.confidentiality != "staff_only":
            if meeting.student_id:
                candidate_ids.update(guardian_ids_for_student(meeting.student_id))
            elif meeting.class_name:
                class_students = User.query.filter(
                    User.role == "student", db.func.lower(User.class_name) == meeting.class_name.casefold(),
                    User.is_active_account.is_(True),
                ).all()
                for student in class_students:
                    candidate_ids.update(guardian_ids_for_student(student.id))
        staff_candidates = User.query.filter(
            User.role.in_(MEETING_ROLES - {"parent"}), User.is_active_account.is_(True),
        ).order_by(User.full_name).all()
        target_class = meeting.class_name or (meeting.student.class_name if meeting.student else None)
        for person in staff_candidates:
            if viewer.role in OFFICE_ROLES or person.role in OFFICE_ROLES:
                candidate_ids.add(person.id)
            elif target_class and target_class.casefold() in staff_classes(person):
                candidate_ids.add(person.id)
        candidate_ids.discard(viewer.id)
        candidate_users = User.query.filter(
            User.id.in_(candidate_ids - active_participant_ids),
            User.is_active_account.is_(True),
        ).order_by(User.full_name).all()
        candidate_options = "".join(
            f"<option value='{person.id}'>{escape(person.full_name)} · {escape(person.role.replace('_', ' ').title())}</option>"
            for person in candidate_users
        )
        if candidate_users:
            participant_add_form = f"<form method='post' class='row g-2 mt-3'><input type='hidden' name='action' value='participant_add'>{_csrf_field()}<div class='col-md-8'><label class='form-label'>Invite another authorized participant</label><select class='form-select' name='user_id' required>{candidate_options}</select></div><div class='col-md-4 d-flex align-items-end'><button class='btn btn-outline-success'>Send invitation</button></div></form>"
    participant_table = (
        "<section class='panel meeting-section'><h2>Invitations and attendance</h2><div class='table-responsive'><table class='table align-middle'><thead><tr><th>Participant</th><th>Role</th><th>Invitation response</th><th>Attendance / representative</th><th>Manage invitation</th></tr></thead><tbody>"
        + ("".join(participant_rows) or "<tr><td colspan='5' class='muted'>No participants are listed.</td></tr>")
        + f"</tbody></table></div>{participant_add_form}</section>"
    )

    agenda_rows = []
    for index, item in enumerate(meeting.agenda_items):
        details = f"<br><small class='muted'>{escape(item.details).replace(chr(10), '<br>')}</small>" if item.details else ""
        item_controls = ""
        if can_manage:
            move_up = f"<form method='post' class='d-inline'><input type='hidden' name='action' value='agenda_move'>{_csrf_field()}<input type='hidden' name='agenda_item_id' value='{item.id}'><input type='hidden' name='direction' value='up'><button class='btn btn-sm btn-outline-secondary' aria-label='Move {escape(item.title)} up' {'disabled' if index == 0 else ''}>↑</button></form>"
            move_down = f"<form method='post' class='d-inline'><input type='hidden' name='action' value='agenda_move'>{_csrf_field()}<input type='hidden' name='agenda_item_id' value='{item.id}'><input type='hidden' name='direction' value='down'><button class='btn btn-sm btn-outline-secondary' aria-label='Move {escape(item.title)} down' {'disabled' if index == len(meeting.agenda_items) - 1 else ''}>↓</button></form>"
            remove_form = f"<form method='post' class='d-inline' onsubmit=\"return confirm('Remove this agenda item?')\"><input type='hidden' name='action' value='agenda_remove'>{_csrf_field()}<input type='hidden' name='agenda_item_id' value='{item.id}'><button class='btn btn-sm btn-outline-danger'>Remove</button></form>"
            item_controls = move_up + move_down + remove_form
        agenda_rows.append(f"<li class='mb-2'><strong>{escape(item.title)}</strong>{details} <span class='ms-2'>{item_controls}</span></li>")
    agenda_form = f"""<form method='post' class='row g-2 mt-2'><input type='hidden' name='action' value='agenda_add'>{_csrf_field()}
<div class='col-md-4'><label class='form-label'>Agenda title</label><input class='form-control' name='title' maxlength='200' required></div>
<div class='col-md-6'><label class='form-label'>Details</label><input class='form-control' name='details' maxlength='5000'></div><div class='col-md-2 d-flex align-items-end'><button class='btn btn-outline-success'>Add item</button></div></form>""" if can_manage else ""
    agenda = "<section class='panel meeting-section'><h2>Agenda</h2><ol>" + ("".join(agenda_rows) or "<li class='muted'>No agenda items recorded.</li>") + f"</ol>{agenda_form}</section>"

    minutes = meeting.minutes
    minutes_panel = ""
    if parent_view:
        if minutes and minutes.status == "approved" and minutes.revisions:
            approved_content = escape(minutes.revisions[-1].content).replace("\n", "<br>")
            minutes_panel = f"<section class='panel meeting-section'><h2>Approved minutes</h2><p class='muted'>Approved by {escape(minutes.approver.full_name if minutes.approver else 'School approver')} · {escape(minutes.approved_at or '')}</p><div>{approved_content}</div></section>"
        elif meeting.status in {"Completed", "Scheduled", "Rescheduled"}:
            minutes_panel = "<section class='panel meeting-section'><h2>Minutes</h2><p class='muted'>Minutes will appear here after school approval.</p></section>"
    elif can_manage:
        minutes_status = escape(minutes.status.replace("_", " ").title()) if minutes else "Not started"
        draft_content = escape(minutes.revisions[-1].content) if minutes and minutes.revisions else ""
        minutes_form = ""
        if not minutes or minutes.status != "pending_approval":
            minutes_form = f"""<form method='post'><input type='hidden' name='action' value='minutes_save'>{_csrf_field()}<label class='form-label'>Minutes draft</label><textarea class='form-control' name='content' rows='8' maxlength='30000' required>{draft_content}</textarea><button class='btn btn-outline-success mt-2'>Save draft</button></form>"""
            if minutes and minutes.status == "draft":
                minutes_form += f"<form method='post' class='mt-2'><input type='hidden' name='action' value='minutes_submit'>{_csrf_field()}<button class='btn btn-accent'>Submit for approval</button></form>"
        decision_form = ""
        if minutes and minutes.status == "pending_approval" and (current_user.role in APPROVER_ROLES and meeting.minutes_approver_id == current_user.id):
            decision_form = f"""<div class='row g-2 mt-2'><form method='post' class='col-md-6'><input type='hidden' name='action' value='minutes_approve'>{_csrf_field()}<label class='form-label'>Approval note</label><input class='form-control' name='note' maxlength='1000'><button class='btn btn-success mt-2'>Approve minutes</button></form><form method='post' class='col-md-6'><input type='hidden' name='action' value='minutes_reject'>{_csrf_field()}<label class='form-label'>Reason for revision</label><input class='form-control' name='note' maxlength='1000' required><button class='btn btn-outline-danger mt-2'>Return for revision</button></form></div>"""
        history = "".join(
            f"<li>v{revision.version} · {escape(revision.event.replace('_', ' ').title())} · {escape(revision.author.full_name if revision.author else 'System')} · {escape(revision.created_at)}{f' · {escape(revision.decision_note)}' if revision.decision_note else ''}</li>"
            for revision in minutes.revisions
        ) if minutes else ""
        minutes_panel = f"<section class='panel meeting-section'><h2>Minutes <span class='pill'>{minutes_status}</span></h2>{minutes_form}{decision_form}{'<h3 class=\"mt-3\">Version history</h3><ol>' + history + '</ol>' if history else ''}</section>"

    action_rows = []
    staff_options = "".join(
        f"<option value='{person.id}'>{escape(person.full_name)} · {escape(person.role.replace('_', ' ').title())}</option>"
        for person in User.query.filter(
            User.role.in_(OFFICE_ROLES | {"teacher", "discipline", "patron", "matron"}),
            User.is_active_account.is_(True),
        ).order_by(User.full_name).all()
    )
    for action in meeting.actions:
        if parent_view and not action.is_parent_visible:
            continue
        comments = "".join(
            f"<li>{escape(comment.author.full_name if comment.author else 'School account')} · {escape(comment.created_at)}: {escape(comment.body)}</li>"
            for comment in action.comments if not parent_view or comment.is_parent_visible
        )
        can_reassign_action = viewer.role in OFFICE_ROLES or viewer.id in {
            action.assigned_by_id, meeting.organizer_id, meeting.chairperson_id, meeting.secretary_id,
        }
        can_update_action = (can_manage or action.assignee_id == viewer.id) and (
            action.status != "Reviewed" or viewer.role in APPROVER_ROLES
        )
        update_form = ""
        if can_update_action and not parent_view:
            allowed_statuses = ["Not Started", "In Progress"]
            if action.assignee_id == viewer.id or viewer.role in APPROVER_ROLES or action.status == "Completed":
                allowed_statuses.append("Completed")
            if action.status == "Overdue":
                allowed_statuses.append("Overdue")
            if viewer.role in APPROVER_ROLES:
                allowed_statuses.append("Reviewed")
            allowed_statuses = list(dict.fromkeys(allowed_statuses))
            status_options = "".join(
                f"<option value='{status}'{' selected' if action.status == status else ''}>{status}</option>"
                for status in allowed_statuses
            )
            deadline_readonly = "" if can_reassign_action else " readonly"
            evidence_readonly = "" if action.assignee_id == viewer.id or viewer.role in APPROVER_ROLES else " readonly"
            review_readonly = "" if viewer.role in APPROVER_ROLES else " readonly"
            assignee_options = staff_options if can_reassign_action else ""
            update_form = f"""<form method='post' class='row g-2 mt-2'><input type='hidden' name='action' value='action_update'>{_csrf_field()}<input type='hidden' name='action_id' value='{action.id}'>
<div class='col-md-3'><label class='form-label'>Status</label><select class='form-select' name='status'>{status_options}</select></div>
<div class='col-md-3'><label class='form-label'>Deadline</label><input class='form-control' type='date' name='deadline' value='{escape(action.deadline)}'{deadline_readonly}></div>
<div class='col-md-3'><label class='form-label'>Assign to</label><select class='form-select' name='assignee_id'><option value='{action.assignee_id}'>{escape(action.assignee.full_name)}</option>{assignee_options}</select></div>
<div class='col-md-3'><label class='form-label'>Completion evidence</label><input class='form-control' name='completion_evidence' value='{escape(action.completion_evidence or '')}' maxlength='5000'{evidence_readonly}></div>
<div class='col-md-8'><label class='form-label'>Review outcome</label><input class='form-control' name='review_outcome' value='{escape(action.review_outcome or '')}' maxlength='1000'{review_readonly}></div>
<div class='col-md-4 d-flex align-items-end'><button class='btn btn-sm btn-outline-success'>Save action</button></div></form>"""
        comment_form = ""
        if can_manage or action.assignee_id == viewer.id or parent_view:
            comment_form = f"<form method='post' class='d-flex gap-2 mt-2'><input type='hidden' name='action' value='action_comment'>{_csrf_field()}<input type='hidden' name='action_id' value='{action.id}'><input class='form-control form-control-sm' name='body' maxlength='2000' required placeholder='Add a follow-up comment'><label class='small'><input type='checkbox' name='is_parent_visible' value='yes'{' checked' if parent_view else ''}> Share with parent</label><button class='btn btn-sm btn-outline-success'>Comment</button></form>"
        action_rows.append(
            f"<article class='border rounded p-3 mb-2'><div class='d-flex justify-content-between gap-2 flex-wrap'><div><strong>{escape(action.title)}</strong><p class='mb-1'>{escape(action.details or '')}</p><small class='muted'>Owner: {escape(action.assignee.full_name)} · Due: {escape(action.deadline)} · Status: {escape(action.status)}</small></div></div>"
            f"{f'<p class=\"small mt-2\">Evidence: {escape(action.completion_evidence)}</p>' if action.completion_evidence and (not parent_view or action.is_parent_visible) else ''}"
            f"{f'<p class=\"small\">Review: {escape(action.review_outcome)}</p>' if action.review_outcome and (not parent_view or action.is_parent_visible) else ''}"
            f"{('<ul class=\"small mt-2\">' + comments + '</ul>') if comments else ''}{update_form}{comment_form}</article>"
        )
    action_form = ""
    if can_manage and minutes and minutes.status == "approved":
        agenda_options = "".join(f"<option value='{item.id}'>{escape(item.title)}</option>" for item in meeting.agenda_items)
        action_form = f"""<details class='mt-3'><summary class='btn btn-sm btn-outline-success'>Assign follow-up from an approved decision</summary>
<form method='post' class='row g-2 mt-2'><input type='hidden' name='action' value='action_create'>{_csrf_field()}
<div class='col-md-4'><label class='form-label'>Action</label><input class='form-control' name='title' maxlength='200' required></div><div class='col-md-3'><label class='form-label'>Owner</label><select class='form-select' name='assignee_id' required>{staff_options}</select></div><div class='col-md-2'><label class='form-label'>Deadline</label><input class='form-control' type='date' name='deadline' required></div><div class='col-md-3'><label class='form-label'>Agenda item</label><select class='form-select' name='agenda_item_id'><option value=''>No linked item</option>{agenda_options}</select></div>
<div class='col-md-8'><label class='form-label'>Details</label><textarea class='form-control' name='details' rows='2'></textarea></div><div class='col-md-4 d-flex align-items-end'><label><input type='checkbox' name='is_parent_visible' value='yes'> Visible to invited parent</label></div><div class='col-12'><button class='btn btn-accent'>Assign action</button></div></form></details>"""
    action_panel = "<section class='panel meeting-section'><h2>Decisions and follow-up actions</h2>" + ("".join(action_rows) or "<p class='muted'>No actions are recorded.</p>") + action_form + "</section>"

    edit_form = ""
    approval_forms = ""
    if can_manage:
        editable_statuses = {meeting.status}
        if meeting.status not in {"Completed", "Cancelled"}:
            editable_statuses.add("Cancelled")
        if meeting.status in {"Scheduled", "Rescheduled"} and viewer.role in OFFICE_ROLES:
            editable_statuses.add("Completed")
        status_options = "".join(
            f"<option value='{status}'{' selected' if status == meeting.status else ''}>{status}</option>"
            for status in sorted(editable_statuses)
        )
        edit_form = f"""<details class='panel meeting-section'><summary><strong>Edit meeting schedule</strong></summary><form method='post' class='row g-2 mt-2'>
<input type='hidden' name='action' value='update'>{_csrf_field()}<div class='col-md-5'><label class='form-label'>Title</label><input class='form-control' name='title' value='{escape(meeting.title)}' maxlength='160' required></div>
<div class='col-md-2'><label class='form-label'>Date</label><input class='form-control' type='date' name='scheduled_date' value='{escape(meeting.scheduled_date)}' required></div>
<div class='col-md-2'><label class='form-label'>Starts</label><input class='form-control' type='time' name='start_time' value='{escape(meeting.start_time)}' required></div>
<div class='col-md-2'><label class='form-label'>Ends</label><input class='form-control' type='time' name='end_time' value='{escape(meeting.end_time)}' required></div>
<div class='col-md-3'><label class='form-label'>Room / location</label><input class='form-control' name='location' value='{escape(meeting.location or '')}' maxlength='160'></div>
<div class='col-md-3'><label class='form-label'>Status</label><select class='form-select' name='status'>{status_options}</select></div>
<div class='col-md-6'><label class='form-label'>Description</label><textarea class='form-control' name='description' rows='2' maxlength='10000'>{escape(meeting.description or '')}</textarea></div>
<div class='col-12'><button class='btn btn-outline-success'>Save changes</button></div></form></details>"""
        if meeting.status == "Pending approval":
            edit_form = f"<form method='post' class='panel meeting-section' onsubmit=\"return confirm('Cancel this meeting submission?')\"><input type='hidden' name='action' value='update'>{_csrf_field()}<input type='hidden' name='status' value='Cancelled'><h2>Cancel submission</h2><p class='muted'>The meeting is locked while it awaits approval. You can cancel the submission.</p><button class='btn btn-outline-danger'>Cancel meeting</button></form>"
        if meeting.status == "Draft":
            approval_forms += f"<form method='post' class='panel meeting-section'><input type='hidden' name='action' value='submit_meeting'>{_csrf_field()}<h2>Approval</h2><p class='muted'>Submit this draft to the designated headteacher or super administrator.</p><button class='btn btn-accent'>Submit for approval</button></form>"
        if meeting.status == "Pending approval" and viewer.role in APPROVER_ROLES and meeting.minutes_approver_id == viewer.id:
            approval_forms += f"""<section class='panel meeting-section'><h2>Meeting approval</h2><div class='row g-2'><form method='post' class='col-md-6'>{_csrf_field()}<input type='hidden' name='action' value='approve_meeting'><label class='form-label'>Approval note</label><input class='form-control' name='note' maxlength='1000'><button class='btn btn-success mt-2'>Approve and send invitations</button></form>
<form method='post' class='col-md-6'>{_csrf_field()}<input type='hidden' name='action' value='reject_meeting'><label class='form-label'>Reason for revision</label><input class='form-control' name='note' maxlength='1000' required><button class='btn btn-outline-danger mt-2'>Return to organizer</button></form></div></section>"""

    can_attend = any(row.user_id == viewer.id for row in meeting.participants)
    related = (
        (response_form if parent_view else "") + panel_details + agenda + participant_table + minutes_panel + action_panel
    )
    return overview.replace("</section>", f"{calendar_status}</section>", 1) + approval_feedback + response_form + edit_form + approval_forms + related


def _report_records():
    records = visible_meetings(current_user)
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    class_name = request.args.get("class_name", "").strip().casefold()
    organizer_value = request.args.get("organizer_id", "").strip()
    meeting_type = request.args.get("type", "").strip()
    status = request.args.get("status", "").strip()
    for value in (date_from, date_to):
        if value:
            try:
                datetime.date.fromisoformat(value)
            except ValueError:
                abort(400, description="Report dates must use YYYY-MM-DD.")
    try:
        organizer_id = int(organizer_value) if organizer_value else None
    except ValueError:
        abort(400, description="Organizer must be a valid account ID.")
    if date_from:
        records = [row for row in records if row.scheduled_date >= date_from]
    if date_to:
        records = [row for row in records if row.scheduled_date <= date_to]
    if class_name:
        records = [row for row in records if (row.class_name or "").casefold() == class_name]
    if organizer_id:
        records = [row for row in records if row.organizer_id == organizer_id]
    if meeting_type in {"general", "class", "parent_teacher", "disciplinary", "emergency", "committee", "panel_discussion"}:
        records = [row for row in records if row.meeting_type == meeting_type]
    if status in MEETING_STATUSES:
        records = [row for row in records if row.status == status]
    return records


def _report_rows(records):
    rows = []
    for meeting in records:
        parents = [row for row in meeting.participants if row.user and row.user.role == "parent"]
        attendance = [row.attendance.status for row in parents if row.attendance]
        rows.append([
            meeting.scheduled_date, meeting.title, meeting.meeting_type.replace("_", " ").title(),
            meeting.class_name or "School-wide", meeting.status,
            len(meeting.participants), len(parents),
            sum(row.invitation_status == "accepted" for row in parents),
            sum(status == "Present" for status in attendance),
            sum(status == "Absent" for status in attendance),
            sum(status == "Late" for status in attendance),
            sum(status == "Excused" for status in attendance),
        ])
    return rows


def _pdf_bytes(title, table_rows):
    """Build a small self-contained text PDF without an optional runtime package."""
    headers = ["Date", "Meeting", "Type", "Class", "Status", "Participants", "Parents invited", "Parents accepted", "Parent present", "Parent absent", "Parent late", "Parent excused"]
    lines = [title, "Generated from authorized meeting records", " | ".join(headers)]
    for row in table_rows:
        lines.append(" | ".join(str(value) for value in row))
    wrapped = []
    for line in lines:
        wrapped.extend(textwrap.wrap(line, width=105, break_long_words=True, break_on_hyphens=False) or [""])
    page_lines = [wrapped[index:index + 48] for index in range(0, len(wrapped), 48)] or [[]]
    objects = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    page_ids = [4 + index * 2 for index in range(len(page_lines))]
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{page_id} 0 R' for page_id in page_ids)}] /Count {len(page_ids)} >>".encode("ascii")
    for page_id, rows in zip(page_ids, page_lines):
        content_id = page_id + 1
        commands = ["BT", "/F1 9 Tf", "45 805 Td", "11 TL"]
        for index, line in enumerate(rows):
            clean = line.encode("latin-1", "replace").decode("latin-1")
            clean = clean.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            if index:
                commands.append("T*")
            commands.append(f"({clean}) Tj")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode("ascii")
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode("ascii") + stream + b"\nendstream")
    output = bytearray(b"%PDF-1.4\n% meeting report\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode("ascii"))
    return bytes(output)


def meeting_reports():
    records = _report_records()
    filters = "&".join(
        f"{key}={escape(request.args.get(key, ''))}"
        for key in ("date_from", "date_to", "class_name", "organizer_id", "type", "status")
        if request.args.get(key)
    )
    query = f"?{filters}" if filters else ""
    organizers = User.query.filter(User.role.in_(OFFICE_ROLES | {"teacher"}),
                                   User.is_active_account.is_(True)).order_by(User.full_name).all()
    organizer_options = "".join(
        f"<option value='{person.id}'{' selected' if request.args.get('organizer_id') == str(person.id) else ''}>{escape(person.full_name)}</option>"
        for person in organizers
    )
    class_names = sorted({row.class_name for row in records if row.class_name}, key=str.casefold)
    class_options = "".join(
        f"<option value='{escape(name)}'{ ' selected' if request.args.get('class_name', '').casefold() == name.casefold() else ''}>{escape(name)}</option>"
        for name in class_names
    )
    rows = _report_rows(records)
    body = f"""<section class='panel'><span class='pill'>AUTHORIZED REPORTS</span><h1 class='mt-2'>Meeting reports</h1>
<p class='muted'>Attendance and participation summaries from meetings in your authorized scope. Exports omit student names, private messages, minutes, and action comments.</p>
<form method='get' class='row g-2 align-items-end'><div class='col-md-2'><label class='form-label'>From</label><input class='form-control' type='date' name='date_from' value='{escape(request.args.get('date_from', ''))}'></div>
<div class='col-md-2'><label class='form-label'>To</label><input class='form-control' type='date' name='date_to' value='{escape(request.args.get('date_to', ''))}'></div>
<div class='col-md-3'><label class='form-label'>Class</label><select class='form-select' name='class_name'><option value=''>All classes</option>{class_options}</select></div>
<div class='col-md-3'><label class='form-label'>Organizer</label><select class='form-select' name='organizer_id'><option value=''>All organizers</option>{organizer_options}</select></div>
<div class='col-md-2'><label class='form-label'>Type</label><select class='form-select' name='type'><option value=''>All</option>{''.join(f"<option value='{name}'{' selected' if request.args.get('type') == name else ''}>{name.replace('_', ' ').title()}</option>" for name in ('general','class','parent_teacher','disciplinary','emergency','committee','panel_discussion'))}</select></div>
<div class='col-md-3'><label class='form-label'>Status</label><select class='form-select' name='status'><option value=''>All statuses</option>{''.join(f"<option value='{escape(name)}{' selected' if request.args.get('status') == name else ''}>{escape(name)}</option>" for name in sorted(MEETING_STATUSES))}</select></div>
<div class='col-md-3'><button class='btn btn-accent'>Apply filters</button> <a class='btn btn-outline-secondary' href='/meetings/reports'>Clear</a></div></form>
<div class='d-flex gap-2 flex-wrap mt-3'><a class='btn btn-outline-success' href='/meetings/reports.xlsx{query}'>Export Excel</a><a class='btn btn-outline-success' href='/meetings/reports.pdf{query}'>Export PDF</a></div></section>
<section class='panel'><div class='table-responsive'><table class='table'><thead><tr><th>Date</th><th>Meeting</th><th>Type</th><th>Class</th><th>Status</th><th>Participants</th><th>Parents invited</th><th>Parents accepted</th><th>Parent present</th><th>Parent absent</th><th>Parent late</th><th>Parent excused</th></tr></thead><tbody>{''.join('<tr>' + ''.join(f'<td>{escape(str(value))}</td>' for value in row) + '</tr>' for row in rows) or "<tr><td colspan='12' class='muted'>No meeting records match these filters.</td></tr>"}</tbody></table></div></section>"""
    return page("Meeting Reports", body, "master_admin")


def meeting_report_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Meeting report"
    sheet.append(["Date", "Meeting", "Type", "Class", "Status", "Participants", "Parents invited", "Parents accepted", "Parent present", "Parent absent", "Parent late", "Parent excused"])
    for row in _report_rows(_report_records()):
        safe_row = list(row)
        for index in (1, 3):
            if isinstance(safe_row[index], str) and safe_row[index].startswith(("=", "+", "-", "@")):
                safe_row[index] = "'" + safe_row[index]
        sheet.append(safe_row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17664C")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 14, "B": 34, "C": 22, "D": 28, "E": 20,
                          "F": 14, "G": 16, "H": 17, "I": 15, "J": 15, "K": 14, "L": 16}.items():
        sheet.column_dimensions[column].width = width
    content = io.BytesIO()
    workbook.save(content)
    content.seek(0)
    return send_file(content, as_attachment=True, download_name="meeting-report.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def meeting_report_pdf():
    data = _pdf_bytes("Parents and School Meeting Report", _report_rows(_report_records()))
    return send_file(io.BytesIO(data), as_attachment=True, download_name="meeting-report.pdf",
                     mimetype="application/pdf")


def meeting_guardians():
    if request.method == "POST":
        if not _csrf_valid():
            abort(400)
        action = request.form.get("action", "")
        try:
            if action == "verify":
                verify_guardian_link(
                    current_user, request.form.get("student_id", type=int),
                    request.form.get("guardian_id", type=int), request.form.get("relationship", ""),
                )
                flash("Guardian relationship verified.", "success")
            elif action == "revoke":
                link = db.session.get(StudentGuardian, request.form.get("link_id", type=int))
                if not link:
                    abort(404)
                revoke_guardian_link(current_user, link)
                flash("Guardian access revoked.", "success")
            else:
                abort(400)
        except PermissionError:
            abort(403)
        except ValueError as error:
            flash(str(error), "danger")
        return redirect(url_for("meeting_guardians"))
    students = User.query.filter_by(role="student", is_active_account=True).order_by(User.class_name, User.full_name).all()
    parents = User.query.filter_by(role="parent", is_active_account=True).order_by(User.full_name).all()
    links = StudentGuardian.query.order_by(StudentGuardian.status, StudentGuardian.student_id, StudentGuardian.guardian_id).all()
    student_options = "".join(f"<option value='{student.id}'>{escape(student.full_name)} · {escape(student.class_name or 'No class')}</option>" for student in students)
    guardian_options = "".join(f"<option value='{parent.id}'>{escape(parent.full_name)}</option>" for parent in parents)
    rows = []
    for link in links:
        action = ""
        if link.status == "verified":
            action = f"<form method='post' onsubmit=\"return confirm('Revoke this guardian relationship?')\"><input type='hidden' name='csrf_token' value='{escape(_csrf_token())}'><input type='hidden' name='action' value='revoke'><input type='hidden' name='link_id' value='{link.id}'><button class='btn btn-sm btn-outline-danger'>Revoke</button></form>"
        rows.append(f"<tr><td>{escape(link.student.full_name if link.student else 'Deleted student')}</td><td>{escape(link.guardian.full_name if link.guardian else 'Deleted guardian')}</td><td>{escape(link.relationship)}</td><td>{escape(link.status.title())}</td><td>{escape(link.verified_at)}</td><td>{escape(link.revoked_at or '')}</td><td>{action}</td></tr>")
    body = f"""<section class='panel'><h1>Verified parent and guardian links</h1><p class='muted'>Only verified links grant access to student-specific meeting details. Revocation immediately blocks that access, including legacy parent links.</p>
<form method='post' class='row g-2 align-items-end'><input type='hidden' name='csrf_token' value='{escape(_csrf_token())}'><input type='hidden' name='action' value='verify'><div class='col-md-4'><label class='form-label'>Student</label><select class='form-select' name='student_id' required>{student_options}</select></div><div class='col-md-4'><label class='form-label'>Parent / guardian account</label><select class='form-select' name='guardian_id' required>{guardian_options}</select></div><div class='col-md-2'><label class='form-label'>Relationship</label><input class='form-control' name='relationship' value='Parent / Guardian' maxlength='80'></div><div class='col-md-2'><button class='btn btn-accent'>Verify link</button></div></form></section>
<section class='panel'><div class='table-responsive'><table class='table'><thead><tr><th>Student</th><th>Guardian</th><th>Relationship</th><th>Status</th><th>Verified</th><th>Revoked</th><th>Action</th></tr></thead><tbody>{''.join(rows) or "<tr><td colspan='7' class='muted'>No guardian links are recorded yet.</td></tr>"}</tbody></table></div></section>"""
    return page("Guardian Verification", body, "master_admin")


def meeting_settings():
    if request.method == "POST":
        if not _csrf_valid():
            abort(400)
        approval_required = "true" if request.form.get("meeting_approval_required") == "yes" else "false"
        try:
            grace_days = int(request.form.get("meeting_overdue_escalation_days", "0"))
        except (TypeError, ValueError):
            flash("Overdue escalation delay must be a whole number from 0 to 30.", "danger")
            return redirect(url_for("meeting_settings"))
        if grace_days < 0 or grace_days > 30:
            flash("Overdue escalation delay must be a whole number from 0 to 30.", "danger")
            return redirect(url_for("meeting_settings"))
        values = {
            "meeting_approval_required": approval_required,
            "meeting_overdue_escalation_days": str(grace_days),
        }
        changed = {}
        for key, value in values.items():
            setting = Setting.query.filter_by(key=key).first()
            old_value = setting.value if setting else None
            if old_value != value:
                changed[key] = {"old": old_value, "new": value}
                if setting:
                    setting.value = value
                else:
                    db.session.add(Setting(key=key, value=value))
        if changed:
            record_meeting_audit(current_user, None, "meeting_configuration_updated", changed)
        db.session.commit()
        flash("Meeting settings saved.", "success")
        return redirect(url_for("meeting_settings"))
    from app.utils.helpers import get_setting
    approval_required = get_setting("meeting_approval_required", "true") != "false"
    grace_days = get_setting("meeting_overdue_escalation_days", "0")
    checked = " checked" if approval_required else ""
    body = f"""<section class='panel'><span class='pill'>SCHOOL SETTINGS</span><h1 class='mt-2'>Meeting workflow settings</h1>
<p class='muted'>These settings apply to new meetings and the daily action escalation command.</p>
<form method='post' class='row g-3'><input type='hidden' name='csrf_token' value='{escape(_csrf_token())}'>
<div class='col-12'><label><input type='checkbox' name='meeting_approval_required' value='yes'{checked}> Require designated headteacher / super administrator approval before staff meetings are published</label></div>
<div class='col-md-4'><label class='form-label' for='meeting-overdue-days'>Days past deadline before action escalation</label><input id='meeting-overdue-days' class='form-control' type='number' name='meeting_overdue_escalation_days' min='0' max='30' value='{escape(grace_days)}' required><small class='muted'>0 escalates on the day after the deadline.</small></div>
<div class='col-12'><button class='btn btn-accent'>Save settings</button></div></form></section>"""
    return page("Meeting Settings", body, "master_admin")


def meeting_audit():
    from app.services.meeting_service import verify_meeting_audit_chain
    valid, events = verify_meeting_audit_chain()
    rows = "".join(
        f"<tr><td>{escape(event.created_at)}</td><td>{escape(event.actor_name)} ({escape(event.actor_role)})</td><td>{escape(event.event_type.replace('_', ' ').title())}</td><td>{escape(str(event.meeting_id or 'Account-level'))}</td><td><code>{escape(event.previous_hash[:14])}…</code></td><td><code>{escape(event.event_hash[:14])}…</code></td></tr>"
        for event in reversed(events[-500:])
    )
    integrity = "Audit chain verified" if valid else "Audit chain integrity check failed"
    tone = "success" if valid else "danger"
    empty_row = "<tr><td colspan='6' class='muted'>No audit events are recorded.</td></tr>"
    body = f"<section class='panel'><h1>Meeting audit trail</h1><div class='alert alert-{tone}' role='status'>{integrity}</div><p class='muted'>The ordered records are HMAC chained and anchored to the latest event hash. Showing up to 500 recent events.</p><div class='table-responsive'><table class='table'><thead><tr><th>Time</th><th>Actor</th><th>Event</th><th>Meeting</th><th>Previous hash</th><th>Event hash</th></tr></thead><tbody>{rows or empty_row}</tbody></table></div></section>"
    return page("Meeting Audit Trail", body, "master_admin")


def register_routes(app):
    app.add_url_rule("/meetings", endpoint="meetings", view_func=permission_for("meetings")(meetings), methods=["GET", "POST"])
    app.add_url_rule("/api/meetings/live-schedule", endpoint="meeting_live_schedule", view_func=permission_for("meetings")(meeting_live_schedule), methods=["GET"])
    app.add_url_rule("/api/meetings/room-bookings", endpoint="room_booking_api", view_func=permission_for("meetings")(room_booking_api), methods=["GET", "POST"])
    app.add_url_rule("/api/meetings/room-booking-qr.svg", endpoint="room_booking_qr", view_func=permission_for("meetings")(room_booking_qr), methods=["GET"])
    app.add_url_rule("/meetings/<int:meeting_id>", endpoint="meeting_detail", view_func=permission_for("meetings")(meeting_detail), methods=["GET", "POST"])
    app.add_url_rule("/meetings/guardians", endpoint="meeting_guardians", view_func=permission_for("meeting_admin")(meeting_guardians), methods=["GET", "POST"])
    app.add_url_rule("/meetings/settings", endpoint="meeting_settings", view_func=permission_for("meeting_settings")(meeting_settings), methods=["GET", "POST"])
    app.add_url_rule("/meetings/reports", endpoint="meeting_reports", view_func=permission_for("meeting_admin")(meeting_reports), methods=["GET"])
    app.add_url_rule("/meetings/reports.xlsx", endpoint="meeting_report_xlsx", view_func=permission_for("meeting_admin")(meeting_report_xlsx), methods=["GET"])
    app.add_url_rule("/meetings/reports.pdf", endpoint="meeting_report_pdf", view_func=permission_for("meeting_admin")(meeting_report_pdf), methods=["GET"])
    app.add_url_rule("/meetings/audit", endpoint="meeting_audit", view_func=permission_for("meeting_audit")(meeting_audit), methods=["GET"])
